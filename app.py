from __future__ import annotations

import io
import csv
import calendar as pycalendar
import json
import re
import os
import secrets
import sqlite3
import smtplib
import unicodedata
from email.message import EmailMessage
from datetime import date, datetime, timedelta
from functools import wraps
from pathlib import Path
from urllib.parse import urlparse

from flask import Flask, Response, abort, flash, g, jsonify, redirect, render_template, request, session, url_for
from werkzeug.security import check_password_hash, generate_password_hash
from manual_search import load_index as load_manual_index, search_manual

BASE_DIR = Path(__file__).resolve().parent
INSTANCE_DIR = BASE_DIR / "instance"
INSTANCE_DIR.mkdir(exist_ok=True)


def load_dotenv(path: Path) -> None:
    """Small .env reader to keep local setup dependency-free."""
    if not path.exists():
        return
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        key, value = key.strip(), value.strip().strip('"').strip("'")
        if key and key not in os.environ:
            os.environ[key] = value


load_dotenv(BASE_DIR / ".env")
DB_PATH = Path(os.getenv("DATABASE_PATH", str(INSTANCE_DIR / "fmabc_eventos.db"))).resolve()
DB_PATH.parent.mkdir(parents=True, exist_ok=True)

app = Flask(__name__)
app.config.update(
    SECRET_KEY=os.getenv("SECRET_KEY") or secrets.token_hex(32),
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.getenv("COOKIE_SECURE", "1" if os.getenv("RENDER") else "0").lower() in {"1", "true", "yes"},
    MAX_CONTENT_LENGTH=8 * 1024 * 1024,
)

STATUS_LABELS = {
    "recebido": "Solicitação recebida",
    "em_triagem": "Em triagem",
    "aguardando_informacoes": "Aguardando informações",
    "aprovado": "Aprovado",
    "em_planejamento": "Em planejamento",
    "realizado": "Realizado (pós-evento pendente)",
    "encerrado": "Encerrado",
    "cancelado": "Cancelado",
}
STATUS_CLASSES = {
    "recebido": "status-blue",
    "em_triagem": "status-purple",
    "aguardando_informacoes": "status-amber",
    "aprovado": "status-green",
    "em_planejamento": "status-teal",
    "realizado": "status-gray",
    "encerrado": "status-green",
    "cancelado": "status-red",
}
EVENT_TYPES = [
    "Congresso / Simpósio", "Aula / Atividade acadêmica", "Cerimônia / Colação de grau",
    "Evento institucional", "Palestra / Mesa-redonda", "Curso / Oficina", "Visita institucional",
    "Reunião", "Inauguração / Lançamento", "Premiação / Homenagem", "Outro",
]
TASK_STATUSES = {
    "pendente": "Não iniciada",
    "em_andamento": "Em andamento",
    "aguardando_terceiro": "Aguardando terceiro",
    "concluida": "Concluída",
    "nao_se_aplica": "Não se aplica",
}
TASK_DONE_STATUSES = {"concluida", "nao_se_aplica"}
SUPPORT_CATALOG_STATUSES = ["Disponível", "Parcial", "Em solicitação", "Não disponível", "Precisa confirmar", "Inativo"]
COMMUNICATION_ROLES = {"acao": "Ação necessária", "conhecimento": "Para conhecimento", "excluir": "Não incluir no comunicado"}
DEPARTMENT_RESPONSE_LABELS = {"aguardando": "Aguardando resposta", "recebido": "Recebido — em análise", "confirmado": "Confirmado pela área", "nao_podera_atender": "Não poderá atender", "precisa_alinhar": "Precisa alinhar", "ciente": "Ciente"}
DEPARTMENT_TASK_STATUSES = {
    "nao_iniciada": "Não iniciada",
    "em_andamento": "Em andamento",
    "aguardando_terceiro": "Aguardando terceiro",
    "bloqueada": "Bloqueada",
    "concluida": "Concluída",
}
DEPARTMENT_TASK_STATUS_LABELS = DEPARTMENT_TASK_STATUSES.copy()
DEPARTMENT_TASK_HISTORY_FIELDS = (
    ("involvement", "Participação da área"),
    ("details", "Providências solicitadas"),
    ("notes", "Observação interna"),
    ("area_responsible", "Responsável da área"),
    ("due_date", "Prazo da providência"),
    ("response_status", "Resposta da área"),
    ("receipt_confirmed_at", "Confirmação de recebimento"),
    ("task_status", "Status da providência"),
    ("pending_reason", "Justificativa da pendência"),
    ("response_notes", "Observações da área"),
)
DEFAULT_DEPARTMENTS = ["TI", "Inspetoria", "Limpeza", "Manutenção", "Imprensa"]
# Opções preservadas do Google Forms original da FMABC (PDF recebido em 20/10/2026).
FORMATS = ["Presencial", "On-line", "Híbrido"]
SUPPORT_OPTIONS = ["Manutenção", "Limpeza", "TI", "Inspetoria", "Outro"]
VENUE_OPTIONS = [
    "Administrativo - salas 5/7 e 6/8 - 60 lugares cada",
    "Auditório da Cardiologia - 40 lugares",
    "Auditório da Dermatologia - 78 lugares",
    "Auditório da Estética - 80 lugares",
    "Auditório da Morfologia - 130 lugares",
    "Auditório da Pós–Graduação - 316 lugares",
    "Auditório da Urologia - 45 lugares",
    "Auditório David Uip - 280 lugares",
    "Sala de videoconferência - 73 lugares",
    "Salão Nobre - 81 lugares",
    "Externo (Fora da FMABC)",
    "Outro",
]
REQUESTED_ITEM_OPTIONS = [
    "Brindes", "Certificado", "Coffee break", "Credencial", "Divulgação prévia",
    "Divulgação posterior", "Equipamentos de áudio visual para apresentação", "Gravação",
    "Inscrição prévia", "Mestre de Cerimônias", "Registro fotográfico",
    "Reserva de Estacionamento", "Transmissão Simultânea", "Outro",
]
AUTHORIZATION_OPTIONS = [
    "Reitoria", "Pró-Reitoria de Extensão", "Pró-Reitoria de Graduação",
    "Pró-Reitoria de Pós-Graduação",
]

# Slots fixos de 30 em 30 minutos para o formulário público de reserva.
# O valor armazenado permanece parseável pelos verificadores legados de horário.
TIME_OPTIONS = [
    (f"{hour:02d}h{minute:02d}", f"{hour:02d}h" if minute == 0 else f"{hour:02d}h30")
    for hour in range(24) for minute in (0, 30)
]
TIME_OPTION_VALUES = {value for value, _label in TIME_OPTIONS}

SCHEMA = """
CREATE TABLE IF NOT EXISTS users (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    email TEXT NOT NULL UNIQUE COLLATE NOCASE,
    password_hash TEXT NOT NULL,
    role TEXT NOT NULL DEFAULT 'staff',
    is_active INTEGER NOT NULL DEFAULT 1,
    created_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS events (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    protocol TEXT NOT NULL UNIQUE,
    title TEXT NOT NULL,
    department TEXT NOT NULL,
    organizer_name TEXT NOT NULL,
    organizer_email TEXT NOT NULL,
    organizer_phone TEXT NOT NULL,
    event_type TEXT NOT NULL,
    event_format TEXT NOT NULL,
    start_date TEXT NOT NULL,
    end_date TEXT,
    start_time TEXT,
    end_time TEXT,
    expected_attendees INTEGER NOT NULL DEFAULT 0,
    venue TEXT,
    audience TEXT,
    description TEXT NOT NULL,
    support_json TEXT NOT NULL DEFAULT '[]',
    event_modes_json TEXT NOT NULL DEFAULT '[]',
    venues_json TEXT NOT NULL DEFAULT '[]',
    requested_items_json TEXT NOT NULL DEFAULT '[]',
    support_other TEXT,
    venue_other TEXT,
    items_other TEXT,
    contact_email TEXT,
    legacy_contact_reference TEXT,
    authorized_by TEXT,
    cost_center TEXT,
    additional_info TEXT,
    event_type_other TEXT,
    external_guests TEXT NOT NULL DEFAULT 'Não informado',
    accessibility TEXT NOT NULL DEFAULT 'Não informado',
    budget_source TEXT,
    partner TEXT,
    status TEXT NOT NULL DEFAULT 'recebido',
    priority TEXT NOT NULL DEFAULT 'Normal',
    internal_owner TEXT,
    internal_notes TEXT,
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS tasks (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    event_id INTEGER NOT NULL REFERENCES events(id) ON DELETE CASCADE,
    title TEXT NOT NULL,
    category TEXT NOT NULL DEFAULT 'Organização',
    due_date TEXT,
    status TEXT NOT NULL DEFAULT 'pendente',
    assigned_to TEXT,
    description TEXT,
    notes TEXT,
    sort_order INTEGER NOT NULL DEFAULT 0,
    created_at TEXT NOT NULL,
    completed_at TEXT
);
CREATE TABLE IF NOT EXISTS task_templates (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    title TEXT NOT NULL,
    category TEXT NOT NULL DEFAULT 'Planejamento',
    description TEXT,
    trigger_type TEXT NOT NULL DEFAULT 'always',
    trigger_value TEXT,
    due_offset INTEGER NOT NULL DEFAULT -7,
    assigned_to TEXT,
    active INTEGER NOT NULL DEFAULT 1,
    sort_order INTEGER NOT NULL DEFAULT 0,
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_task_templates_active ON task_templates(active, sort_order);
CREATE TABLE IF NOT EXISTS support_catalog (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL,
    category TEXT NOT NULL DEFAULT 'Material',
    quantity INTEGER,
    location TEXT,
    status TEXT NOT NULL DEFAULT 'Precisa confirmar',
    notes TEXT,
    active INTEGER NOT NULL DEFAULT 1,
    updated_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_support_catalog_active ON support_catalog(active, category, name);
CREATE INDEX IF NOT EXISTS idx_events_status ON events(status);
CREATE INDEX IF NOT EXISTS idx_events_start_date ON events(start_date);
CREATE INDEX IF NOT EXISTS idx_tasks_event_id ON tasks(event_id);
CREATE TABLE IF NOT EXISTS departments (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL UNIQUE COLLATE NOCASE,
    recipient_email TEXT,
    active INTEGER NOT NULL DEFAULT 1,
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_departments_active ON departments(active, name);
CREATE TABLE IF NOT EXISTS event_department_updates (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    event_id INTEGER NOT NULL REFERENCES events(id) ON DELETE CASCADE,
    department_id INTEGER NOT NULL REFERENCES departments(id),
    involvement TEXT NOT NULL DEFAULT 'conhecimento',
    details TEXT,
    notes TEXT,
    updated_by TEXT,
    updated_at TEXT NOT NULL,
    response_status TEXT NOT NULL DEFAULT 'aguardando',
    response_notes TEXT,
    response_by TEXT,
    response_token TEXT,
    responded_at TEXT,
    receipt_confirmed_at TEXT,
    receipt_confirmed_by TEXT,
    area_responsible TEXT,
    due_date TEXT,
    task_status TEXT NOT NULL DEFAULT 'nao_iniciada',
    task_updated_at TEXT,
    task_updated_by TEXT,
    pending_reason TEXT,
    UNIQUE(event_id, department_id)
);
CREATE INDEX IF NOT EXISTS idx_event_department_updates_event ON event_department_updates(event_id);
CREATE TABLE IF NOT EXISTS department_task_history (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    update_id INTEGER NOT NULL REFERENCES event_department_updates(id) ON DELETE CASCADE,
    event_id INTEGER NOT NULL REFERENCES events(id) ON DELETE CASCADE,
    department_id INTEGER NOT NULL REFERENCES departments(id),
    changed_at TEXT NOT NULL,
    changed_by TEXT,
    source TEXT NOT NULL DEFAULT 'interno',
    action TEXT NOT NULL,
    change_summary TEXT NOT NULL,
    before_json TEXT,
    after_json TEXT
);
CREATE INDEX IF NOT EXISTS idx_department_task_history_update ON department_task_history(update_id, changed_at DESC);
CREATE INDEX IF NOT EXISTS idx_department_task_history_event_department ON department_task_history(event_id, department_id, changed_at DESC);
CREATE TABLE IF NOT EXISTS communication_batches (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    department_id INTEGER NOT NULL REFERENCES departments(id),
    period_year INTEGER NOT NULL,
    period_month INTEGER NOT NULL,
    subject TEXT NOT NULL,
    body TEXT NOT NULL,
    recipient_email TEXT,
    status TEXT NOT NULL DEFAULT 'rascunho',
    created_by TEXT,
    created_at TEXT NOT NULL,
    sent_at TEXT,
    sent_by TEXT,
    send_method TEXT,
    error_message TEXT
);
CREATE INDEX IF NOT EXISTS idx_communication_batches_period ON communication_batches(period_year, period_month, department_id, created_at);
CREATE TABLE IF NOT EXISTS communication_batch_events (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    batch_id INTEGER NOT NULL REFERENCES communication_batches(id) ON DELETE CASCADE,
    event_id INTEGER NOT NULL REFERENCES events(id),
    event_title TEXT NOT NULL,
    event_date TEXT NOT NULL,
    event_time TEXT,
    event_venue TEXT,
    involvement TEXT NOT NULL,
    details TEXT,
    UNIQUE(batch_id, event_id)
);
CREATE INDEX IF NOT EXISTS idx_communication_batch_events_batch ON communication_batch_events(batch_id);
CREATE TABLE IF NOT EXISTS event_closures (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    event_id INTEGER NOT NULL UNIQUE REFERENCES events(id) ON DELETE CASCADE,
    actual_attendees INTEGER,
    summary TEXT NOT NULL,
    positives TEXT,
    issues TEXT,
    improvements TEXT,
    follow_up_actions TEXT,
    final_observations TEXT,
    closed_by TEXT,
    closed_at TEXT NOT NULL,
    updated_by TEXT,
    updated_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_event_closures_closed_at ON event_closures(closed_at);
CREATE TABLE IF NOT EXISTS event_evaluations (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    event_id INTEGER NOT NULL UNIQUE REFERENCES events(id) ON DELETE CASCADE,
    token TEXT NOT NULL UNIQUE,
    requester_email TEXT NOT NULL,
    rating_overall INTEGER,
    rating_sector INTEGER,
    rating_communication INTEGER,
    rating_responsiveness INTEGER,
    rating_needs_met INTEGER,
    positive_feedback TEXT,
    improvement_suggestions TEXT,
    additional_comments TEXT,
    created_at TEXT NOT NULL,
    invitation_sent_at TEXT,
    invitation_method TEXT,
    responded_at TEXT
);
CREATE INDEX IF NOT EXISTS idx_event_evaluations_responded_at ON event_evaluations(responded_at);
CREATE TABLE IF NOT EXISTS event_program_items (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    event_id INTEGER NOT NULL REFERENCES events(id) ON DELETE CASCADE,
    start_time TEXT NOT NULL,
    end_time TEXT,
    activity TEXT NOT NULL,
    location TEXT,
    responsible TEXT,
    notes TEXT,
    sort_order INTEGER NOT NULL DEFAULT 0,
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_event_program_items_event ON event_program_items(event_id, start_time, sort_order);
CREATE TABLE IF NOT EXISTS event_budget_items (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    event_id INTEGER NOT NULL REFERENCES events(id) ON DELETE CASCADE,
    item_name TEXT NOT NULL,
    category TEXT NOT NULL DEFAULT 'Outros',
    supplier TEXT,
    estimated_amount REAL,
    actual_amount REAL,
    status TEXT NOT NULL DEFAULT 'previsto',
    notes TEXT,
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_event_budget_items_event ON event_budget_items(event_id, category, item_name);
CREATE TABLE IF NOT EXISTS event_incidents (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    event_id INTEGER NOT NULL REFERENCES events(id) ON DELETE CASCADE,
    incident_date TEXT,
    category TEXT NOT NULL DEFAULT 'Outro',
    title TEXT NOT NULL,
    impact TEXT,
    resolution TEXT,
    preventive_action TEXT,
    responsible TEXT,
    status TEXT NOT NULL DEFAULT 'aberta',
    recorded_by TEXT,
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_event_incidents_event ON event_incidents(event_id, status, incident_date);
"""


def get_db() -> sqlite3.Connection:
    if "db" not in g:
        conn = sqlite3.connect(DB_PATH)
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA foreign_keys = ON")
        conn.execute("PRAGMA journal_mode = WAL")
        g.db = conn
    return g.db


@app.teardown_appcontext
def close_db(_error=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def now_iso() -> str:
    return datetime.now().replace(microsecond=0).isoformat(sep=" ")


def init_db() -> None:
    db = get_db()
    db.executescript(SCHEMA)
    # Atualiza bancos locais já criados em versões anteriores, sem apagar registros.
    existing_columns = {row["name"] for row in db.execute("PRAGMA table_info(events)").fetchall()}
    migrations = {
        "event_modes_json": "TEXT NOT NULL DEFAULT '[]'",
        "venues_json": "TEXT NOT NULL DEFAULT '[]'",
        "requested_items_json": "TEXT NOT NULL DEFAULT '[]'",
        "support_other": "TEXT",
        "venue_other": "TEXT",
        "items_other": "TEXT",
        "contact_email": "TEXT",
        "legacy_contact_reference": "TEXT",
        "authorized_by": "TEXT",
        "cost_center": "TEXT",
        "additional_info": "TEXT",
        "event_type_other": "TEXT",
    }
    for column, sql_type in migrations.items():
        if column not in existing_columns:
            db.execute(f"ALTER TABLE events ADD COLUMN {column} {sql_type}")
    user_columns = {row["name"] for row in db.execute("PRAGMA table_info(users)").fetchall()}
    if "is_active" not in user_columns:
        db.execute("ALTER TABLE users ADD COLUMN is_active INTEGER NOT NULL DEFAULT 1")
    task_columns = {row["name"] for row in db.execute("PRAGMA table_info(tasks)").fetchall()}
    task_migrations = {"description": "TEXT", "notes": "TEXT"}
    for column, sql_type in task_migrations.items():
        if column not in task_columns:
            db.execute(f"ALTER TABLE tasks ADD COLUMN {column} {sql_type}")
    department_update_columns = {row["name"] for row in db.execute("PRAGMA table_info(event_department_updates)").fetchall()}
    department_update_migrations = {
        "response_status": "TEXT NOT NULL DEFAULT 'aguardando'", "response_notes": "TEXT",
        "response_by": "TEXT", "response_token": "TEXT", "responded_at": "TEXT",
        "receipt_confirmed_at": "TEXT", "receipt_confirmed_by": "TEXT",
        "area_responsible": "TEXT", "due_date": "TEXT",
        "task_status": "TEXT NOT NULL DEFAULT 'nao_iniciada'", "task_updated_at": "TEXT",
        "task_updated_by": "TEXT", "pending_reason": "TEXT",
    }
    for column, sql_type in department_update_migrations.items():
        if column not in department_update_columns:
            db.execute(f"ALTER TABLE event_department_updates ADD COLUMN {column} {sql_type}")
    db.execute("CREATE UNIQUE INDEX IF NOT EXISTS idx_event_department_response_token ON event_department_updates(response_token)")
    # Vínculos de providência antigos também recebem um token individual para retorno seguro.
    legacy_area_rows = db.execute("SELECT id FROM event_department_updates WHERE involvement='acao' AND (response_token IS NULL OR response_token='')").fetchall()
    for legacy_row in legacy_area_rows:
        db.execute("UPDATE event_department_updates SET response_token=? WHERE id=?", (secrets.token_urlsafe(32), legacy_row["id"]))
    db.commit()
    seed_task_templates(db)
    seed_support_catalog(db)
    seed_departments(db)
    current = db.execute("SELECT id FROM users LIMIT 1").fetchone()
    if not current:
        email = os.getenv("ADMIN_EMAIL", "eventos@fmabc.br").strip().lower()
        password = os.getenv("ADMIN_PASSWORD")
        generated = not password
        if generated:
            password = secrets.token_urlsafe(14)
        elif len(password) < 6:
            raise RuntimeError("ADMIN_PASSWORD deve ter pelo menos 6 caracteres para criar a conta administrativa inicial.")
        db.execute(
            "INSERT INTO users(email, password_hash, role, created_at) VALUES (?, ?, 'admin', ?)",
            (email, generate_password_hash(password), now_iso()),
        )
        db.commit()
        print("\n[FMABC Eventos] Usuário administrativo inicial criado.")
        print(f"E-mail: {email}")
        print(f"Senha inicial: {password}")
        print("Defina ADMIN_PASSWORD e SECRET_KEY no .env antes de usar em ambiente institucional.\n")


def seed_departments(db: sqlite3.Connection) -> None:
    """Create initial area list; recipient addresses remain blank until confirmed by FMABC."""
    for name in DEFAULT_DEPARTMENTS:
        db.execute(
            "INSERT OR IGNORE INTO departments(name, recipient_email, active, created_at, updated_at) VALUES (?, NULL, 1, ?, ?)",
            (name, now_iso(), now_iso()),
        )
    db.commit()


def smtp_settings() -> dict:
    return {
        "host": os.getenv("SMTP_HOST", "").strip(),
        "port": int(os.getenv("SMTP_PORT", "587") or 587),
        "user": os.getenv("SMTP_USER", "").strip(),
        "password": os.getenv("SMTP_PASSWORD", ""),
        "from_email": os.getenv("SMTP_FROM", "").strip(),
        "use_tls": os.getenv("SMTP_USE_TLS", "1").lower() in {"1", "true", "yes"},
    }


def smtp_is_configured() -> bool:
    cfg = smtp_settings()
    return bool(cfg["host"] and cfg["from_email"])


def build_monthly_message(department, year: int, month: int, events: list[dict]) -> tuple[str, str]:
    month_names = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"]
    subject = f"Agenda de eventos FMABC — {month_names[month - 1].capitalize()} de {year} — {department['name']}"
    lines = [
        "Prezados,", "", 
        f"Encaminhamos a programação de eventos da FMABC prevista para {month_names[month - 1]} de {year}, para organização e conhecimento da área {department['name']}.",
        "As indicações de envolvimento e as observações abaixo refletem os registros atuais do Setor de Eventos. Favor sinalizar eventuais ajustes.", ""
    ]
    for idx, item in enumerate(events, start=1):
        try:
            d = date.fromisoformat(item["start_date"]).strftime("%d/%m/%Y")
        except (TypeError, ValueError):
            d = item["start_date"] or "Data a confirmar"
        lines.append(f"{idx}. {item['title']}")
        time_text = (item.get("start_time") or "").strip()
        end_text = (item.get("end_time") or "").strip()
        if time_text and end_text and end_text not in time_text:
            time_text = f"{time_text} às {end_text}"
        lines.append(f"   Data: {d}" + (f" | Horário: {time_text}" if time_text else ""))
        lines.append(f"   Local: {item['venue'] or 'A confirmar'}")
        responsible = (item.get("organizer_name") or "").strip()
        if responsible:
            lines.append(f"   Responsável registrado: {responsible}")
        legacy_contact = (item.get("legacy_contact_reference") or "").strip()
        if legacy_contact:
            lines.append(f"   Contato/referência registrada na agenda: {legacy_contact}")
        else:
            contacts = []
            if (item.get("organizer_phone") or "").strip():
                contacts.append("Telefone/WhatsApp: " + item["organizer_phone"].strip())
            if (item.get("contact_email") or "").strip():
                contacts.append("E-mail de contato: " + item["contact_email"].strip())
            for contact_line in contacts:
                lines.append("   " + contact_line)
        lines.append(f"   Participação da área: {COMMUNICATION_ROLES.get(item['involvement'], item['involvement'])}")
        if item.get("area_responsible"):
            lines.append(f"   Responsável da área: {item['area_responsible']}")
        if item.get("due_date"):
            try:
                due_text = date.fromisoformat(item['due_date']).strftime("%d/%m/%Y")
            except (TypeError, ValueError):
                due_text = item['due_date']
            lines.append(f"   Prazo da providência: {due_text}")
        if item.get("details"):
            lines.append(f"   Providências solicitadas: {item['details']}")
        if item.get("response_url"):
            lines.append(f"   Confirme o recebimento, informe o responsável e atualize a situação da providência: {item['response_url']}")
        elif item.get("involvement") == "acao":
            lines.append("   Ação necessária: confirme o recebimento, indique o responsável da área e atualize a situação da providência.")
        lines.append("")
    lines += ["Pedimos que as áreas envolvidas organizem as providências indicadas e comuniquem o Setor de Eventos sobre possíveis impedimentos ou alterações.", "", "Atenciosamente,", "Setor de Eventos", "Centro Universitário FMABC"]
    return subject, "\n".join(lines)


@app.before_request
def csrf_protection():
    if request.method in {"POST", "PUT", "PATCH", "DELETE"}:
        supplied = request.form.get("_csrf", "")
        expected = session.get("_csrf", "")
        if not supplied or not expected or not secrets.compare_digest(supplied, expected):
            abort(400, description="A sessão do formulário expirou. Atualize a página e tente novamente.")


@app.context_processor
def inject_template_globals():
    if "_csrf" not in session:
        session["_csrf"] = secrets.token_urlsafe(32)
    return {
        "csrf_token": session["_csrf"],
        "status_labels": STATUS_LABELS,
        "status_classes": STATUS_CLASSES,
        "event_types": EVENT_TYPES,
        "task_statuses": TASK_STATUSES,
        "support_catalog_statuses": SUPPORT_CATALOG_STATUSES,
        "formats": FORMATS,
        "support_options": SUPPORT_OPTIONS,
        "venue_options": VENUE_OPTIONS,
        "requested_item_options": REQUESTED_ITEM_OPTIONS,
        "authorization_options": AUTHORIZATION_OPTIONS,
        "time_options": TIME_OPTIONS,
        "communication_roles": COMMUNICATION_ROLES,
        "department_response_labels": DEPARTMENT_RESPONSE_LABELS,
        "department_task_statuses": DEPARTMENT_TASK_STATUSES,
        "department_task_status_labels": DEPARTMENT_TASK_STATUS_LABELS,
        "today": date.today().isoformat(),
    }


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        user_id = session.get("user_id")
        if not user_id:
            return redirect(url_for("login", next=request.path))
        # Re-check the account on every protected request so blocked/deleted accounts
        # lose access immediately on their next navigation, even with an old session.
        user = get_db().execute(
            "SELECT id, email, role, is_active FROM users WHERE id = ?", (user_id,)
        ).fetchone()
        if not user or not user["is_active"]:
            session.clear()
            flash("Esta conta está bloqueada ou não está mais cadastrada. Entre em contato com um administrador.", "error")
            return redirect(url_for("login"))
        # Apply role/email changes immediately to sessions that are already open.
        session["user_email"] = user["email"]
        session["user_role"] = user["role"]
        return view(*args, **kwargs)
    return wrapped


def get_event_or_404(event_id: int):
    event = get_db().execute("SELECT * FROM events WHERE id = ?", (event_id,)).fetchone()
    if not event:
        abort(404)
    return event


def json_list(event, field: str) -> list[str]:
    try:
        data = json.loads(event[field] or "[]")
        return data if isinstance(data, list) else []
    except (json.JSONDecodeError, TypeError, KeyError):
        return []


def event_support(event) -> list[str]:
    return json_list(event, "support_json")


def event_modes(event) -> list[str]:
    values = json_list(event, "event_modes_json")
    return values or [x.strip() for x in (event["event_format"] or "").split(",") if x.strip()]


def event_venues(event) -> list[str]:
    values = json_list(event, "venues_json")
    return values or [x.strip() for x in (event["venue"] or "").split("; ") if x.strip()]


def event_requested_items(event) -> list[str]:
    return json_list(event, "requested_items_json")


def seed_task_templates(db: sqlite3.Connection) -> None:
    """Seed a first editable checklist model based on the FMABC manual checklist."""
    if db.execute("SELECT 1 FROM task_templates LIMIT 1").fetchone():
        return
    templates = []
    def t(title, category, description, trigger_type="always", trigger_value=None, offset=-7, owner=None):
        templates.append((title, category, description, trigger_type, trigger_value, offset, owner, len(templates) + 1))

    # General tasks — always created for approved events.
    t("Conferir formulário de solicitação e dados do evento", "Planejamento", "Validar data, horário, objetivo, setor solicitante, público e contatos.", offset=-45)
    t("Conferir autorização do evento e assinaturas necessárias", "Administrativo", "Verificar a documentação institucional e registrar pendências.", offset=-40)
    t("Confirmar reserva do local e capacidade", "Infraestrutura", "Conferir espaços selecionados, capacidade, acessibilidade e disponibilidade.", offset=-40)
    t("Consolidar programação do evento", "Planejamento", "Reunir horários, atividades e responsáveis pela programação.", offset=-30)
    t("Avaliar necessidade de reunião inicial com a área demandante", "Planejamento", "Registrar se a reunião será necessária e, se aplicável, agendar com o solicitante.", offset=-30)
    t("Confirmar programação e consolidar informações para o roteiro", "Planejamento", "Reunir horários, participantes, responsáveis e informações que dependem da área solicitante.", offset=-20)
    t("Montar cronograma operacional e designar responsáveis", "Planejamento", "Definir as etapas, os prazos e os responsáveis internos/externos.", offset=-25)
    t("Confirmar equipamentos e suporte técnico necessários", "Infraestrutura", "Relacionar audiovisual, internet, computador, microfones e apoio necessário.", offset=-14)
    t("Conferir materiais institucionais e sinalização", "Materiais", "Verificar totem, banners, kits e materiais que serão utilizados no espaço.", offset=-7)
    t("Realizar alinhamento final com as áreas envolvidas", "Operação", "Confirmar que solicitações e dependências estão atendidas.", offset=-3)
    t("Conferir montagem e preparação do espaço", "Infraestrutura", "Verificar mesas, cadeiras, púlpito, banners, credenciamento e organização do ambiente.", offset=-1)
    t("Organizar recepção dos convidados e equipe de apoio", "Recepção", "Definir ponto de recepção, responsáveis, fluxo de chegada e apoio aos convidados.", offset=-3)
    t("Executar o evento e registrar ocorrências", "Dia do evento", "Acompanhar a operação, registrar intercorrências e providências tomadas.", offset=0)
    t("Organizar fotos e materiais de cobertura após o evento", "Pós-evento", "Reunir registros e encaminhar materiais aos responsáveis pela publicação.", offset=1)
    t("Avaliar pendências e registrar feedback do evento", "Pós-evento", "Realizar reunião de feedback quando necessária e registrar aprendizados.", offset=7)

    # Tasks triggered by the original form's support selections.
    t("Confirmar solicitação de manutenção e vistoria do espaço", "Manutenção", "Verificar necessidades de manutenção e registrar confirmação do atendimento.", "support", "Manutenção", -7)
    t("Alinhar limpeza antes, durante e após o evento", "Limpeza", "Confirmar horários, escopo e responsável pelo atendimento.", "support", "Limpeza", -3)
    t("Alinhar apoio da inspetoria, acessos e circulação", "Operação", "Confirmar acessos, circulação, apoio durante o evento e horários da equipe.", "support", "Inspetoria", -3)
    t("Testar computador, projeção, som, microfones e internet", "TI / Audiovisual", "Realizar teste integrado e confirmar disponibilidade de apoio técnico.", "support", "TI", -2)
    t("Planejar e aprovar artes e divulgação prévia", "Comunicação", "Alinhar conteúdo, canais, público e aprovação da área solicitante.", "item", "Divulgação prévia", -20)
    t("Organizar divulgação posterior do evento", "Comunicação", "Selecionar registros e definir texto/canais para a cobertura posterior.", "item", "Divulgação posterior", 1)
    t("Configurar inscrições prévias e validar lista de participantes", "Inscrições", "Definir plataforma, período de inscrição, confirmação e exportação da lista.", "item", "Inscrição prévia", -14)
    t("Preparar lista de presença ou QR Code", "Credenciamento", "Definir método, campos e processo para controlar a presença.", "item", "Inscrição prévia", -3)
    t("Definir modelo e critérios para emissão de certificados", "Certificados", "Validar critérios de presença e fluxo de emissão para público e palestrantes.", "item", "Certificado", -7)
    t("Confirmar cardápio, quantidade, horário e serviço de coffee break", "Fornecedores", "Validar público previsto, local de montagem, toalhas/apoios e fornecedor.", "item", "Coffee break", -5)
    t("Preparar arte, dados e quantidade de credenciais/crachás", "Credenciamento", "Confirmar nome, identificação, arte e prazo para produção.", "item", "Credencial", -7)
    t("Preparar etiquetas de identificação e conferir nomes", "Credenciamento", "Confirmar quem receberá etiquetas, grafia dos nomes e quantidade necessária.", "item", "Credencial", -3)
    t("Confirmar toalhas de TNT e mesas de apoio para o coffee break", "Materiais", "Conferir quantidade de toalhas e local de montagem do serviço.", "item", "Coffee break", -3)
    t("Definir quantidade, personalização e entrega dos brindes", "Materiais", "Verificar estoque, personalização, prazo e disponibilidade antes de confirmar.", "item", "Brindes", -10)
    t("Confirmar mestre de cerimônias e revisar roteiro", "Cerimonial", "Validar roteiro, ordem das falas e se haverá mesa solene.", "item", "Mestre de Cerimônias", -7)
    t("Confirmar registro fotográfico e momentos prioritários", "Comunicação", "Alinhar responsável, cobertura e autorização para registros.", "item", "Registro fotográfico", -3)
    t("Confirmar gravação, permissões e entrega do material", "Audiovisual", "Definir formato de gravação, autorizações e destino dos arquivos.", "item", "Gravação", -5)
    t("Testar plataforma, conexão e transmissão simultânea", "Transmissão", "Realizar ensaio com participantes remotos e testar áudio, vídeo e internet.", "item", "Transmissão Simultânea", -2)
    t("Realizar teste técnico do ambiente remoto e do acesso dos participantes", "Transmissão", "Testar link, acesso de convidados, áudio, vídeo, conexão e apoio a participantes remotos.", "mode", "On-line", -2)
    t("Realizar teste técnico do ambiente remoto e do acesso dos participantes", "Transmissão", "Testar link, acesso de convidados, áudio, vídeo, conexão e apoio a participantes remotos.", "mode", "Híbrido", -2)
    t("Confirmar quantidade, nomes e logística do estacionamento", "Logística", "Validar vagas necessárias, convidados autorizados, local e liberação.", "item", "Reserva de Estacionamento", -3)
    t("Confirmar mesas, cadeiras, púlpito, banners e montagem especial", "Infraestrutura", "Alinhar layout do espaço, materiais de montagem e horário de preparação.", "item", "Equipamentos de áudio visual para apresentação", -3)

    # Additional tasks from the manual checklist, triggered by event category.
    t("Confirmar palestrantes/convidados e receber minibiografias e fotos", "Convidados", "Acompanhar confirmações, biografias, fotos e necessidades dos participantes.", "category", "Congresso / Simpósio", -14)
    t("Confirmar palestrantes/convidados e receber minibiografias e fotos", "Convidados", "Acompanhar confirmações, biografias, fotos e necessidades dos participantes.", "category", "Palestra / Mesa-redonda", -14)
    t("Validar roteiro cerimonial, mesa solene e ordem das falas", "Cerimonial", "Confirmar autoridades, precedência, roteiro, mestre de cerimônias e tempo de cada fala.", "category", "Cerimônia / Colação de grau", -14)
    t("Organizar recepção institucional e sequência de falas", "Cerimonial", "Confirmar equipe de recepção, convidados institucionais e roteiro de abertura.", "category", "Inauguração / Lançamento", -7)
    t("Preparar premiação, homenagens e sequência de entrega", "Cerimonial", "Validar nomes, ordem, materiais, certificados/prêmios e responsáveis pela entrega.", "category", "Premiação / Homenagem", -7)
    t("Alinhar recepção de convidados e equipe de apoio", "Recepção", "Confirmar ponto de recepção, responsáveis e fluxo de chegada.", "category", "Visita institucional", -3)

    db.executemany(
        "INSERT INTO task_templates(title, category, description, trigger_type, trigger_value, due_offset, assigned_to, active, sort_order, created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?, ?, 1, ?, ?, ?)",
        [(title, category, desc, trigger, value, offset, owner, order, now_iso(), now_iso())
         for title, category, desc, trigger, value, offset, owner, order in templates],
    )


def seed_support_catalog(db: sqlite3.Connection) -> None:
    """Seed editable support inventory notes from the manual checklist; ambiguous stock is flagged."""
    if db.execute("SELECT 1 FROM support_catalog LIMIT 1").fetchone():
        return
    records = [
        ("Faixas", "Comunicação", None, None, "Precisa confirmar", "Item citado no checklist manual; confirmar existência, quantidade e responsável."),
        ("Banners", "Comunicação", None, None, "Precisa confirmar", "Confirmar modelos disponíveis, quantidade e localização."),
        ("Crachás", "Credenciamento", None, None, "Em solicitação", "Checklist manual contém a anotação 'Pedido no compras'. Confirmar situação atual."),
        ("Tendas", "Infraestrutura", None, None, "Precisa confirmar", "Confirmar disponibilidade, quantidade, local de armazenamento e manutenção."),
        ("Porta-banners", "Materiais", None, "IC / Inspetoria / Fisiologia — confirmar", "Precisa confirmar", "O checklist manual menciona 50 IC, 32 Inspetoria e 20 Fisiologia, além de '102 disponível' e '110 unidades'. Valores divergentes: fazer inventário antes de informar quantidade disponível."),
        ("Brindes", "Materiais", None, None, "Precisa confirmar", "Checklist manual inclui a observação 'Não temos disponível'; confirmar estoque atual antes de prometer entrega."),
        ("Mesas extras", "Infraestrutura", None, None, "Precisa confirmar", "Confirmar quantidade disponível e reserva com a área responsável."),
        ("Toalhas de TNT para coffee break", "Materiais", None, None, "Precisa confirmar", "Verificar quantidade, cor e local de armazenamento; estimar pelo público e formato do serviço."),
        ("Totem / material institucional", "Comunicação", None, None, "Precisa confirmar", "Confirmar materiais que podem ser utilizados no evento e quem fará o transporte."),
        ("Estacionamento para convidados", "Logística", None, None, "Precisa confirmar", "Confirmar quantidade de vagas, local, nomes autorizados e fluxo de liberação."),
        ("Coffee break — mesas/apoios", "Infraestrutura", None, None, "Precisa confirmar", "Definir local de montagem, horário de acesso e materiais necessários."),
        ("Restaurante FSA", "Fornecedores", None, None, "Precisa confirmar", "O checklist manual cita este restaurante; confirmar contato, disponibilidade, condições e prazo antes de indicar ao solicitante."),
        ("Restaurante Quintal do ABC", "Fornecedores", None, None, "Precisa confirmar", "O checklist manual cita este restaurante e sugere aviso prévio; confirmar contato, disponibilidade, condições e prazo."),
    ]
    db.executemany(
        "INSERT INTO support_catalog(name, category, quantity, location, status, notes, active, updated_at) VALUES (?, ?, ?, ?, ?, ?, 1, ?)",
        [(*record, now_iso()) for record in records],
    )


def add_checklist(event) -> int:
    """Create an initial checklist from configurable templates and the event form selections."""
    db = get_db()
    if db.execute("SELECT 1 FROM tasks WHERE event_id = ? LIMIT 1", (event["id"],)).fetchone():
        return 0
    try:
        event_day = date.fromisoformat(event["start_date"])
    except (TypeError, ValueError):
        event_day = date.today()

    services = event_support(event)
    requested = event_requested_items(event)
    modes = event_modes(event)
    category = event["event_type"] or ""
    selected_support = set(services)
    selected_items = set(requested)
    selected_modes = set(modes)

    templates = db.execute(
        "SELECT * FROM task_templates WHERE active = 1 ORDER BY sort_order, id"
    ).fetchall()
    inserted = 0
    added_keys = set()
    for order, template in enumerate(templates, start=1):
        trigger = template["trigger_type"]
        value = (template["trigger_value"] or "").strip()
        include = (
            trigger == "always"
            or (trigger == "support" and value in selected_support)
            or (trigger == "item" and value in selected_items)
            or (trigger == "mode" and value in selected_modes)
            or (trigger == "category" and value == category)
        )
        if not include:
            continue
        duplicate_key = (template["title"].strip().casefold(), template["category"].strip().casefold())
        if duplicate_key in added_keys:
            continue
        added_keys.add(duplicate_key)
        due_date = (event_day + timedelta(days=int(template["due_offset"] or 0))).isoformat()
        db.execute(
            """INSERT INTO tasks(event_id, title, category, due_date, status, assigned_to, description, notes, sort_order, created_at)
               VALUES (?, ?, ?, ?, 'pendente', ?, ?, NULL, ?, ?)""",
            (event["id"], template["title"], template["category"], due_date,
             template["assigned_to"], template["description"], inserted + 1, now_iso()),
        )
        inserted += 1
    db.commit()
    return inserted


def normalize_date(value: str | None) -> str | None:
    if not value:
        return None
    try:
        return date.fromisoformat(value).isoformat()
    except ValueError:
        return None


def create_event_from_form(form):
    required = [
        "organizer_email", "start_date", "title", "start_time", "end_time", "department", "organizer_name",
        "audience", "expected_attendees", "description", "organizer_phone", "contact_email",
        "authorized_by", "cost_center", "event_category",
    ]
    missing = [field for field in required if not (form.get(field) or "").strip()]
    if missing:
        return None, "Preencha todos os campos obrigatórios."

    start_date = normalize_date(form.get("start_date"))
    if not start_date:
        return None, "Informe uma data válida para o evento."
    try:
        attendees = int(form.get("expected_attendees") or "")
        if attendees < 0:
            raise ValueError
    except ValueError:
        return None, "Informe um número válido de participantes."

    email = form.get("organizer_email", "").strip().lower()
    contact_email = form.get("contact_email", "").strip().lower()
    for value in (email, contact_email):
        if "@" not in value or "." not in value.split("@")[-1]:
            return None, "Informe um e-mail válido nos campos de contato."

    modes = [item for item in form.getlist("event_modes") if item in FORMATS]
    services = [item for item in form.getlist("support") if item in SUPPORT_OPTIONS]
    venues = [item for item in form.getlist("venues") if item in VENUE_OPTIONS]
    requested_items = [item for item in form.getlist("requested_items") if item in REQUESTED_ITEM_OPTIONS]
    authorization = form.get("authorized_by", "").strip()
    event_type = form.get("event_category", "").strip()
    event_type_other = form.get("event_category_other", "").strip() or None
    if event_type not in EVENT_TYPES:
        return None, "Selecione a categoria do evento."
    if event_type == "Outro" and not event_type_other:
        return None, "Especifique a categoria do evento selecionando 'Outro'."
    if not modes:
        return None, "Selecione pelo menos uma opção na pergunta 7: tipo do evento."
    if not services:
        return None, "Selecione pelo menos uma opção na pergunta 8: serviços de apoio."
    if not venues:
        return None, "Selecione pelo menos uma opção na pergunta 9: local do evento."
    if not requested_items:
        return None, "Selecione pelo menos uma opção na pergunta 13: itens necessários."
    if authorization not in AUTHORIZATION_OPTIONS:
        return None, "Selecione quem autorizou o evento."

    support_other = form.get("support_other", "").strip() or None
    venue_other = form.get("venue_other", "").strip() or None
    items_other = form.get("items_other", "").strip() or None
    if "Outro" in services and not support_other:
        return None, "Especifique o outro serviço de apoio selecionado."
    if "Outro" in venues and not venue_other:
        return None, "Especifique o outro local selecionado."
    if "Outro" in requested_items and not items_other:
        return None, "Especifique o outro item selecionado."

    # Os horários do formulário são opções fixas de 30 em 30 minutos.
    start_time = form.get("start_time", "").strip()
    end_time = form.get("end_time", "").strip()
    if start_time not in TIME_OPTION_VALUES or end_time not in TIME_OPTION_VALUES:
        return None, "Selecione os horários de início e término nas opções disponíveis."
    start_minutes = _time_to_minutes(start_time)
    end_minutes = _time_to_minutes(end_time)
    if start_minutes is None or end_minutes is None or end_minutes <= start_minutes:
        return None, "O horário de término precisa ser posterior ao horário de início."

    # A disponibilidade é verificada no servidor imediatamente antes do cadastro.
    # Isso impede envio mesmo se a verificação dinâmica do navegador não for executada.
    db = get_db()
    # Reserva o lock de escrita antes da verificação para evitar dois envios concorrentes
    # aceitarem o mesmo espaço/horário ao mesmo tempo.
    db.execute("BEGIN IMMEDIATE")
    blockers = _reservation_conflict_details(
        start_date=start_date,
        start_time=start_time,
        end_time=end_time,
        selected_venues=venues,
        venue_other=venue_other,
        db=db,
    )
    if blockers:
        db.rollback()
        return None, _reservation_conflict_message(blockers)

    # Campos legados mantidos para compatibilidade com os registros anteriores do MVP.
    formatted_modes = ", ".join(modes)
    formatted_venues = "; ".join([v for v in venues if v != "Outro"] + ([venue_other] if "Outro" in venues and venue_other else []))
    protocol = f"FMABC-{date.today().year}-{secrets.token_hex(3).upper()}"
    created = now_iso()
    cursor = db.execute(
        """INSERT INTO events(
        protocol, title, department, organizer_name, organizer_email, organizer_phone,
        event_type, event_format, start_date, end_date, start_time, end_time,
        expected_attendees, venue, audience, description, support_json, event_modes_json,
        venues_json, requested_items_json, support_other, venue_other, items_other,
        contact_email, authorized_by, cost_center, additional_info, event_type_other,
        external_guests, accessibility, budget_source, partner, status, created_at, updated_at
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'recebido', ?, ?)""",
        (protocol, form.get("title", "").strip(), form.get("department", "").strip(),
         form.get("organizer_name", "").strip(), email, form.get("organizer_phone", "").strip(),
         event_type, formatted_modes, start_date, start_date, start_time, end_time,
         attendees, formatted_venues or None, form.get("audience", "").strip(), form.get("description", "").strip(),
         json.dumps(services, ensure_ascii=False), json.dumps(modes, ensure_ascii=False),
         json.dumps(venues, ensure_ascii=False), json.dumps(requested_items, ensure_ascii=False),
         support_other, venue_other, items_other, contact_email, authorization,
         form.get("cost_center", "").strip(), form.get("additional_info", "").strip() or None,
         event_type_other, "Não informado", "Não informado", None, None, created, created),
    )
    db.commit()
    return {"id": cursor.lastrowid, "protocol": protocol}, None


@app.get("/solicitar")
def request_event():
    return render_template("request_form.html", submitted=False)


@app.post("/solicitar")
def submit_event_request():
    result, error = create_event_from_form(request.form)
    if error:
        flash(error, "error")
        return render_template("request_form.html", submitted=False, form=request.form), 400
    return render_template("request_form.html", submitted=True, protocol=result["protocol"])


@app.get("/api/agenda/verificar-reserva")
def api_check_reservation_availability():
    """Read-only availability check used by the public request form."""
    start_date = normalize_date(request.args.get("start_date"))
    start_time = (request.args.get("start_time") or "").strip()
    end_time = (request.args.get("end_time") or "").strip()
    selected_venues = [v for v in request.args.getlist("venues") if v in VENUE_OPTIONS]
    venue_other = (request.args.get("venue_other") or "").strip() or None

    if not start_date or not start_time or not end_time or not selected_venues:
        return jsonify({"ready": False, "conflict": False, "message": "Informe a data, o horário de início, o horário de término e pelo menos um local para verificar a disponibilidade."})
    if start_time not in TIME_OPTION_VALUES or end_time not in TIME_OPTION_VALUES:
        return jsonify({"ready": True, "valid_time": False, "conflict": False,
                        "message": "Selecione os horários de início e término nas opções disponíveis, em intervalos de 30 minutos."})
    start_minutes = _time_to_minutes(start_time)
    end_minutes = _time_to_minutes(end_time)
    if start_minutes is None or end_minutes is None or end_minutes <= start_minutes:
        return jsonify({"ready": True, "valid_time": False, "conflict": False,
                        "message": "O horário de término precisa ser posterior ao horário de início."})
    if "Outro" in selected_venues and not venue_other:
        return jsonify({"ready": False, "conflict": False, "message": "Informe o nome do local selecionando 'Outro' para verificar a disponibilidade."})

    blockers = _reservation_conflict_details(
        start_date=start_date,
        start_time=start_time,
        end_time=end_time,
        selected_venues=selected_venues,
        venue_other=venue_other,
        db=get_db(),
    )
    if blockers:
        return jsonify({"ready": True, "valid_time": True, "conflict": True,
                        "message": _reservation_conflict_message(blockers),
                        "venues": sorted({item["venue"] for item in blockers})})
    if not _selected_venue_labels(selected_venues, venue_other):
        return jsonify({"ready": True, "valid_time": True, "conflict": False,
                        "message": "O local informado é externo ou não é gerenciado na agenda de espaços da FMABC. Confirme a disponibilidade diretamente com o responsável."})
    return jsonify({"ready": True, "valid_time": True, "conflict": False,
                    "message": "Não foi encontrada outra reserva no mesmo local, data e horário. O evento ainda dependerá de análise e aprovação do Setor de Eventos."})


@app.route("/acompanhar", methods=["GET", "POST"])
def track_request():
    result = None
    attempted = False
    if request.method == "POST":
        attempted = True
        protocol = request.form.get("protocol", "").strip().upper()
        email = request.form.get("email", "").strip().lower()
        result = get_db().execute(
            "SELECT protocol, title, start_date, status, created_at, updated_at, internal_notes FROM events WHERE protocol = ? AND lower(organizer_email) = ?",
            (protocol, email),
        ).fetchone()
    return render_template("tracking.html", result=result, attempted=attempted)


def _default_login_destination(role: str | None) -> str:
    """Staff users start in triage; administrators start at the executive dashboard."""
    return url_for("triage_queue") if role == "staff" else url_for("dashboard")


@app.route("/login", methods=["GET", "POST"])
def login():
    if session.get("user_id"):
        existing = get_db().execute(
            "SELECT id, email, role, is_active FROM users WHERE id = ?", (session.get("user_id"),)
        ).fetchone()
        if existing and existing["is_active"]:
            session["user_email"] = existing["email"]
            session["user_role"] = existing["role"]
            return redirect(_default_login_destination(existing["role"]))
        session.clear()
    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        user = get_db().execute("SELECT * FROM users WHERE email = ?", (email,)).fetchone()
        if user and check_password_hash(user["password_hash"], password):
            if not user["is_active"]:
                flash("Esta conta está bloqueada. Entre em contato com o administrador.", "error")
                return render_template("login.html"), 403
            session.clear()
            session["user_id"] = user["id"]
            session["user_email"] = user["email"]
            session["user_role"] = user["role"]
            session["_csrf"] = secrets.token_urlsafe(32)
            # Staff accounts always start in triage, including when a stale `next`
            # parameter points to another page. Administrators may continue to the
            # requested internal page when the URL is local and safe.
            if user["role"] == "staff":
                return redirect(url_for("triage_queue"))
            next_url = request.args.get("next", "")
            if next_url.startswith("/") and not next_url.startswith("//"):
                return redirect(next_url)
            return redirect(_default_login_destination(user["role"]))
        flash("E-mail ou senha inválidos.", "error")
    return render_template("login.html")

@app.post("/logout")
@login_required
def logout():
    session.clear()
    flash("Você saiu do sistema.", "success")
    return redirect(url_for("login"))


@app.get("/manual")
def manual_center():
    """Consulta pública do manual: recupera somente trechos do documento institucional."""
    question = (request.args.get("q") or "").strip()[:350]
    results = search_manual(question) if question else []
    if not question:
        answer_message = "Faça uma pergunta para localizar as orientações correspondentes no manual."
    elif results:
        answer_message = "Encontrei os seguintes trechos no Manual de Cerimonial e Eventos da FMABC. As citações abaixo são passagens do documento original; consulte a página indicada para ler o contexto completo."
    else:
        answer_message = "Não encontrei trechos suficientes no manual para responder com segurança a essa pergunta. Tente usar termos mais próximos dos utilizados no documento ou consulte o manual completo."
    index_info = load_manual_index()
    return render_template(
        "manual.html",
        question=question,
        manual_results=results,
        answer_message=answer_message,
        manual_title=index_info.get("title", "Manual de Cerimonial e Eventos FMABC"),
        manual_year=index_info.get("version_year", 2023),
        manual_pages=index_info.get("pdf_pages", 67),
        manual_pdf_url=url_for("static", filename="docs/manual_cerimonial_eventos_fmabc_2023.pdf"),
    )


def _event_venue_keys(event) -> set[str]:
    values = []
    try:
        values.extend(event_venues(event))
    except Exception:
        pass
    try:
        scalar = event["venue"] or ""
    except (KeyError, TypeError, IndexError):
        scalar = ""
    if scalar:
        values.extend([part.strip() for part in re.split(r"[;|]", scalar) if part.strip()])
    try:
        venue_other = event["venue_other"] or ""
    except (KeyError, TypeError, IndexError):
        venue_other = ""
    if venue_other:
        values.append(venue_other)
    keys = set()
    for value in values:
        clean = _normalize_venue_key(value)
        if clean:
            keys.add(clean)
    return keys


def _time_values_to_minutes(value):
    """Return all valid HH:MM/HHhMM times found in a free-text time field."""
    if not value:
        return []
    results = []
    pattern = r"(?<!\d)(\d{1,2})\s*[:hH]\s*(\d{2})(?!\d)"
    for match in re.finditer(pattern, str(value).strip()):
        hour, minute = int(match.group(1)), int(match.group(2))
        if hour <= 23 and minute <= 59:
            results.append(hour * 60 + minute)
    return results


def _time_to_minutes(value):
    values = _time_values_to_minutes(value)
    return values[0] if values else None


def _time_interval_from_values(start_time, end_time=None):
    parsed_values = _time_values_to_minutes(start_time)
    if not parsed_values:
        return None, None
    start = parsed_values[0]
    end = _time_to_minutes(end_time)
    if end is None and len(parsed_values) > 1:
        end = parsed_values[1]
    if end is not None and end <= start:
        end = None
    return start, end


def _normalize_venue_key(value):
    raw = unicodedata.normalize("NFKD", str(value or ""))
    clean = "".join(char for char in raw if not unicodedata.combining(char))
    clean = clean.casefold().replace("–", "-").replace("—", "-")
    clean = re.sub(r"\s+", " ", clean).strip()
    clean = re.sub(r"\s*-\s*\d+\s*lugares(?:\s+cada)?$", "", clean)
    clean = re.sub(r"^administrativo\s*-\s*", "", clean)
    if clean in {"outro", "externo (fora da fmabc)", "externo", "a confirmar", "nao informado"}:
        return None
    aliases = {
        "auditorio cardiologia": "auditorio da cardiologia",
        "auditorio dermatologia": "auditorio da dermatologia",
        "auditorio estetica": "auditorio da estetica",
        "auditorio morfologia": "auditorio da morfologia",
        "auditorio urologia": "auditorio da urologia",
        "auditorio pos-graduacao": "auditorio da pos-graduacao",
        "salas 5/7 e 6/8": "salas 5/7 e 6/8",
    }
    clean = aliases.get(clean, clean)
    return clean or None


def _selected_venue_labels(selected_venues, venue_other=None):
    labels = []
    for value in selected_venues or []:
        if value == "Outro":
            if venue_other:
                labels.append(venue_other.strip())
        elif value != "Externo (Fora da FMABC)":
            labels.append(value)
    # Deduplicate labels without changing display order.
    result, seen = [], set()
    for label in labels:
        key = _normalize_venue_key(label)
        if key and key not in seen:
            seen.add(key)
            result.append((key, label.strip()))
    return result


def _booking_intervals_overlap(request_start, request_end, existing_start, existing_end):
    if existing_start is None:
        # Existing booking has no parseable time. Caller treats it as an unresolved conflict.
        return False
    if request_end is not None and existing_end is not None:
        return request_start < existing_end and existing_start < request_end
    if request_end is not None:
        return request_start <= existing_start < request_end
    if existing_end is not None:
        return existing_start <= request_start < existing_end
    # When neither duration is known, block identical start times only.
    return request_start == existing_start


def _reservation_conflict_details(start_date, start_time, selected_venues, venue_other=None, db=None, end_time=None):
    """Find same-date room conflicts before requests are recorded; canceled events are ignored."""
    db = db or get_db()
    venue_pairs = _selected_venue_labels(selected_venues, venue_other)
    if not venue_pairs:
        return []  # External venues or unknown locations are not managed as FMABC rooms.
    request_start, request_end = _time_interval_from_values(start_time, end_time)
    if request_start is None:
        return []
    if request_end is not None and request_end <= request_start:
        return []
    selected_map = dict(venue_pairs)
    selected_keys = set(selected_map)

    rows = db.execute(
        """SELECT * FROM events
           WHERE lower(COALESCE(status, '')) <> 'cancelado'
             AND start_date <= ?
             AND COALESCE(NULLIF(end_date, ''), start_date) >= ?
           ORDER BY start_date, start_time""",
        (start_date, start_date),
    ).fetchall()
    blockers = []
    seen = set()
    requested_day = date.fromisoformat(start_date)
    for row in rows:
        event = dict(row)
        existing_keys = _event_venue_keys(row)
        shared = selected_keys & existing_keys
        if not shared:
            continue
        try:
            event_start = date.fromisoformat(event.get("start_date") or "")
            event_end = date.fromisoformat(event.get("end_date") or event.get("start_date") or "")
        except ValueError:
            continue

        # A multi-day booking occupies the room throughout intermediate days; block conservatively.
        all_day_overlap = event_start <= requested_day <= event_end and event_start != event_end
        existing_start, existing_end = _time_interval_from_values(event.get("start_time"), event.get("end_time"))
        unknown_time = existing_start is None
        unknown_duration = existing_start is not None and existing_end is None
        overlaps = all_day_overlap
        unresolved = False
        if not all_day_overlap:
            if unknown_time:
                # Sem horário inicial: a agenda não consegue descartar uma sobreposição.
                unresolved = True
            elif unknown_duration:
                # Há horário inicial, mas não há término. Se o evento começou antes
                # do fim do intervalo solicitado, a reserva pode continuar ocupando o espaço.
                unresolved = (existing_start == request_start) if request_end is None else (existing_start < request_end)
            else:
                overlaps = _booking_intervals_overlap(
                    request_start, request_end, existing_start, existing_end
                )
        if not overlaps and not unresolved:
            continue
        for key in shared:
            venue_label = selected_map.get(key, key)
            marker = (event.get("id"), key, "horario_indefinido" if unresolved else "conflito")
            if marker in seen:
                continue
            seen.add(marker)
            blockers.append({
                "venue": venue_label,
                "kind": "horario_indefinido" if unresolved else "conflito",
                "event_id": event.get("id"),
            })
    return blockers


def _reservation_conflict_message(blockers):
    confirmed = sorted({item["venue"] for item in blockers if item["kind"] == "conflito"})
    unresolved = sorted({item["venue"] for item in blockers if item["kind"] == "horario_indefinido"})
    parts = []
    if confirmed:
        spaces = ", ".join(f'“{name}”' for name in confirmed)
        noun = "O espaço" if len(confirmed) == 1 else "Os espaços"
        verb = "já está reservado" if len(confirmed) == 1 else "já estão reservados"
        parts.append(f"{noun} {spaces} {verb} nesta data e horário para outro evento. Escolha outra data, horário ou local antes de enviar o formulário.")
    if unresolved:
        spaces = ", ".join(f'“{name}”' for name in unresolved)
        parts.append(f"Já existe uma reserva ou solicitação para {spaces} nesta data, mas o horário de início ou término está ausente/incompleto. Para evitar sobreposição, confirme a disponibilidade com o Setor de Eventos ou escolha outro local/data antes de enviar.")
    return " ".join(parts) or "Existe um possível conflito de reserva. Escolha outro local, data ou horário."


def _event_date_bounds(event):
    try:
        start = date.fromisoformat(event["start_date"])
        end_raw = event["end_date"] or event["start_date"]
        end = date.fromisoformat(end_raw)
        return start, end
    except (TypeError, ValueError, KeyError, IndexError):
        return None, None


def _find_schedule_conflicts(events):
    """Detect overlapping bookings of the same known venue; uncertain times are flagged for review."""
    records = [dict(e) for e in events]
    found, seen = [], set()
    for i, first in enumerate(records):
        if first.get("status") == "cancelado":
            continue
        a_start, a_end = _event_date_bounds(first)
        a_venues = _event_venue_keys(first)
        if not a_start or not a_end or not a_venues:
            continue
        for second in records[i + 1:]:
            if second.get("status") == "cancelado" or first.get("id") == second.get("id"):
                continue
            b_start, b_end = _event_date_bounds(second)
            if not b_start or not b_end:
                continue
            overlap_start, overlap_end = max(a_start, b_start), min(a_end, b_end)
            if overlap_start > overlap_end:
                continue
            venues = a_venues & _event_venue_keys(second)
            if not venues:
                continue
            a_st, a_en = _time_to_minutes(first.get("start_time")), _time_to_minutes(first.get("end_time"))
            b_st, b_en = _time_to_minutes(second.get("start_time")), _time_to_minutes(second.get("end_time"))
            single_day = a_start == a_end == b_start == b_end
            if single_day and None not in (a_st, a_en, b_st, b_en):
                if not (a_st < b_en and b_st < a_en):
                    continue
                reason = "Conflito de horário confirmado"
            elif single_day:
                reason = "Possível conflito: duração ou horário final não informado"
            else:
                reason = "Possível conflito: datas sobrepostas no mesmo espaço"
            venue = sorted(venues)[0]
            key = tuple(sorted((first.get("id"), second.get("id")))) + (overlap_start.isoformat(), venue)
            if key in seen:
                continue
            seen.add(key)
            found.append({"event_a_id": first.get("id"), "event_a_title": first.get("title"), "event_b_id": second.get("id"), "event_b_title": second.get("title"), "date": overlap_start.isoformat(), "venue": venue, "reason": reason})
    return sorted(found, key=lambda x: (x["date"], x["venue"], x["event_a_title"] or ""))


def _critical_task_rows(limit=12):
    today = date.today().isoformat()
    soon = (date.today() + timedelta(days=3)).isoformat()
    rows = get_db().execute("""SELECT t.*, e.title event_title, e.start_date event_date, e.id parent_event_id
        FROM tasks t JOIN events e ON e.id=t.event_id
        WHERE t.status NOT IN ('concluida','nao_se_aplica') AND t.due_date IS NOT NULL AND t.due_date<>'' AND t.due_date<=?
          AND e.status NOT IN ('cancelado','encerrado','realizado')
        ORDER BY t.due_date, e.start_date LIMIT ?""", (soon, limit)).fetchall()
    out=[]
    for row in rows:
        item=dict(row); item["overdue"] = item["due_date"] < today
        item["days_until_due"] = (date.fromisoformat(item["due_date"])-date.today()).days
        out.append(item)
    return out


def _department_task_snapshot(row):
    if row is None:
        return {}
    keys = {"involvement", "details", "notes", "area_responsible", "due_date", "response_status",
            "receipt_confirmed_at", "receipt_confirmed_by", "task_status", "task_updated_at",
            "task_updated_by", "pending_reason", "response_notes", "response_by", "responded_at"}
    available = set(row.keys()) if hasattr(row, "keys") else set(row)
    return {key: row[key] for key in keys if key in available}


def _record_department_task_history(db, update_id, event_id, department_id, before, after,
                                    changed_by, source, action):
    before = before or {}
    after = after or {}
    display_changes = []
    for key, label in DEPARTMENT_TASK_HISTORY_FIELDS:
        old, new = before.get(key), after.get(key)
        if old != new:
            old_text = "vazio" if old in (None, "") else str(old)
            new_text = "vazio" if new in (None, "") else str(new)
            if key == "involvement":
                old_text = COMMUNICATION_ROLES.get(old, old_text)
                new_text = COMMUNICATION_ROLES.get(new, new_text)
            elif key == "response_status":
                old_text = DEPARTMENT_RESPONSE_LABELS.get(old, old_text)
                new_text = DEPARTMENT_RESPONSE_LABELS.get(new, new_text)
            elif key == "task_status":
                old_text = DEPARTMENT_TASK_STATUS_LABELS.get(old, old_text)
                new_text = DEPARTMENT_TASK_STATUS_LABELS.get(new, new_text)
            display_changes.append(f"{label}: {old_text} → {new_text}")
    if not display_changes and before:
        return
    summary = "; ".join(display_changes) if display_changes else "Providência criada."
    db.execute("""INSERT INTO department_task_history
        (update_id,event_id,department_id,changed_at,changed_by,source,action,change_summary,before_json,after_json)
        VALUES (?,?,?,?,?,?,?,?,?,?)""",
        (update_id,event_id,department_id,now_iso(),(changed_by or "")[:160] or None,source,action,summary,
         json.dumps(before,ensure_ascii=False,default=str),json.dumps(after,ensure_ascii=False,default=str)))


def _pending_department_confirmations(limit=12):
    return get_db().execute("""SELECT u.*, e.title event_title, e.start_date event_date, e.id parent_event_id,
        d.name department_name, d.recipient_email FROM event_department_updates u
        JOIN events e ON e.id=u.event_id JOIN departments d ON d.id=u.department_id
        WHERE u.involvement='acao' AND (u.task_status<>'concluida'
           OR u.response_status IN ('aguardando','recebido','precisa_alinhar','nao_podera_atender'))
        AND u.response_token IS NOT NULL AND e.start_date>=? AND e.status IN ('aprovado','em_planejamento')
        ORDER BY CASE WHEN u.task_status='bloqueada' THEN 0 WHEN u.response_status='nao_podera_atender' THEN 1
                      WHEN u.response_status IN ('aguardando','recebido','precisa_alinhar') THEN 2 ELSE 3 END,
        COALESCE(u.due_date,e.start_date),e.start_date,d.name LIMIT ?""", (date.today().isoformat(), limit)).fetchall()


@app.get("/agenda")
@login_required
def institutional_calendar():
    today = date.today()
    try:
        year = int(request.args.get("year", str(today.year))); month = int(request.args.get("month", str(today.month)))
        if not 2000 <= year <= 2100 or not 1 <= month <= 12: raise ValueError
    except ValueError:
        year, month = today.year, today.month
    month_start = date(year, month, 1)
    month_end = date(year + (month == 12), 1 if month == 12 else month + 1, 1) - timedelta(days=1)
    events = [dict(r) for r in get_db().execute("""SELECT * FROM events WHERE status<>'cancelado'
        AND start_date<=? AND COALESCE(NULLIF(end_date,''),start_date)>=? ORDER BY start_date,start_time,title""", (month_end.isoformat(), month_start.isoformat())).fetchall()]
    by_day={}
    for item in events:
        start, end = _event_date_bounds(item)
        if not start or not end: continue
        cursor=max(start,month_start); last=min(end,month_end)
        while cursor<=last:
            by_day.setdefault(cursor.isoformat(), []).append(item); cursor += timedelta(days=1)
    weeks=pycalendar.Calendar(firstweekday=0).monthdatescalendar(year,month)
    conflicts=_find_schedule_conflicts(get_db().execute("SELECT * FROM events WHERE status<>'cancelado' AND start_date<=? AND COALESCE(NULLIF(end_date,''),start_date)>=? ORDER BY start_date,start_time", ((month_end+timedelta(days=1)).isoformat(),(month_start-timedelta(days=1)).isoformat())).fetchall())
    return render_template("agenda.html", year=year, month=month, today=today, years=list(range(max(2024,year-2),min(2100,year+5)+1)), weeks=weeks, events_by_day=by_day, conflicts=conflicts, critical_tasks=_critical_task_rows(10), pending_confirmations=_pending_department_confirmations(10), month_names=["Janeiro","Fevereiro","Março","Abril","Maio","Junho","Julho","Agosto","Setembro","Outubro","Novembro","Dezembro"])


@app.route("/confirmar-area/<token>", methods=["GET", "POST"])
def department_response(token):
    if not re.fullmatch(r"[A-Za-z0-9_-]{24,80}", token or ""):
        abort(404)
    db = get_db()
    row = db.execute("""SELECT u.*,e.title event_title,e.protocol,e.start_date,e.start_time,e.end_time,e.venue,
        d.name department_name FROM event_department_updates u
        JOIN events e ON e.id=u.event_id JOIN departments d ON d.id=u.department_id
        WHERE u.response_token=?""", (token,)).fetchone()
    if not row:
        abort(404)

    if request.method == "POST":
        response = (request.form.get("response_status") or "").strip()
        task_status = (request.form.get("task_status") or "nao_iniciada").strip()
        response_by = (request.form.get("response_by") or "").strip()[:120]
        area_responsible = (request.form.get("area_responsible") or "").strip()[:120] or None
        response_notes = (request.form.get("response_notes") or "").strip()[:1500] or None
        pending_reason = (request.form.get("pending_reason") or "").strip()[:1500] or None
        receipt_confirmed = request.form.get("receipt_confirmed") == "yes"

        if not receipt_confirmed:
            flash("Confirme o recebimento da solicitação para registrar a resposta.", "error")
            return redirect(url_for("department_response", token=token))
        if not response_by:
            flash("Informe o nome de quem está respondendo pela área.", "error")
            return redirect(url_for("department_response", token=token))
        if response not in {"recebido", "confirmado", "nao_podera_atender", "precisa_alinhar", "ciente"}:
            flash("Selecione uma resposta válida.", "error")
            return redirect(url_for("department_response", token=token))
        if task_status not in DEPARTMENT_TASK_STATUSES:
            flash("Selecione um status de providência válido.", "error")
            return redirect(url_for("department_response", token=token))
        needs_reason = task_status in {"aguardando_terceiro", "bloqueada"} or response in {"precisa_alinhar", "nao_podera_atender"}
        if needs_reason and not pending_reason:
            flash("Explique a pendência ou o impedimento para que o Setor de Eventos possa acompanhar.", "error")
            return redirect(url_for("department_response", token=token))

        old = _department_task_snapshot(row)
        stamp = now_iso()
        db.execute("""UPDATE event_department_updates SET response_status=?,response_notes=?,response_by=?,responded_at=?,
            receipt_confirmed_at=COALESCE(receipt_confirmed_at,?),receipt_confirmed_by=COALESCE(receipt_confirmed_by,?),
            area_responsible=?,task_status=?,task_updated_at=?,task_updated_by=?,pending_reason=? WHERE id=?""",
            (response,response_notes,response_by,stamp,stamp,response_by,area_responsible,task_status,stamp,response_by,pending_reason,row["id"]))
        updated = db.execute("SELECT * FROM event_department_updates WHERE id=?", (row["id"],)).fetchone()
        _record_department_task_history(db,row["id"],row["event_id"],row["department_id"],old,
            _department_task_snapshot(updated),response_by,"Área de apoio","Resposta/atualização da providência")
        db.commit()
        flash("Recebimento e situação da providência registrados. Obrigado por atualizar o Setor de Eventos.", "success")
        return redirect(url_for("department_response",token=token,saved="1"))

    return render_template("department_response.html", update=row, saved=request.args.get("saved")=="1",
        area_tasks=[], task_statuses=DEPARTMENT_TASK_STATUSES)


@app.route("/eventos/<int:event_id>/roteiro", methods=["GET","POST"])
@login_required
def event_run_sheet(event_id):
    event=get_event_or_404(event_id); db=get_db()
    if request.method=="POST":
        action=request.form.get("action","add"); item_id=request.form.get("item_id",type=int)
        if action=="delete":
            db.execute("DELETE FROM event_program_items WHERE id=? AND event_id=?",(item_id,event_id)); db.commit(); flash("Atividade removida do roteiro.","success")
            return redirect(url_for("event_run_sheet",event_id=event_id))
        start=(request.form.get("start_time") or "").strip(); end=(request.form.get("end_time") or "").strip() or None; activity=(request.form.get("activity") or "").strip()
        if not re.fullmatch(r"\d{2}:\d{2}",start) or (end and not re.fullmatch(r"\d{2}:\d{2}",end)) or not activity:
            flash("Informe a atividade e horários válidos (HH:MM).","error"); return redirect(url_for("event_run_sheet",event_id=event_id))
        if end and _time_to_minutes(end)<=_time_to_minutes(start):
            flash("O horário de término deve ser posterior ao início.","error"); return redirect(url_for("event_run_sheet",event_id=event_id))
        vals=(start,end,activity[:240],(request.form.get("location") or "").strip()[:180] or None,(request.form.get("responsible") or "").strip()[:180] or None,(request.form.get("notes") or "").strip()[:1500] or None,now_iso())
        if action=="edit" and item_id:
            db.execute("UPDATE event_program_items SET start_time=?,end_time=?,activity=?,location=?,responsible=?,notes=?,updated_at=? WHERE id=? AND event_id=?",vals+(item_id,event_id)); flash("Atividade do roteiro atualizada.","success")
        else:
            order=db.execute("SELECT COALESCE(MAX(sort_order),0)+1 FROM event_program_items WHERE event_id=?",(event_id,)).fetchone()[0]
            db.execute("INSERT INTO event_program_items(event_id,start_time,end_time,activity,location,responsible,notes,sort_order,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?,?,?)",(event_id,start,end,activity[:240],vals[3],vals[4],vals[5],order,now_iso(),now_iso())); flash("Atividade adicionada ao roteiro operacional.","success")
        db.commit(); return redirect(url_for("event_run_sheet",event_id=event_id))
    items=db.execute("SELECT * FROM event_program_items WHERE event_id=? ORDER BY start_time,sort_order,id",(event_id,)).fetchall()
    return render_template("event_run_sheet.html",event=event,items=items)


def _parse_money(raw):
    value=(raw or "").strip().replace("R$","").replace(" ","")
    if not value: return None
    if "," in value: value=value.replace(".","").replace(",",".")
    try: amount=float(value)
    except ValueError: raise ValueError("Informe valor monetário válido, por exemplo 1500,00.")
    if amount<0 or amount>1_000_000_000: raise ValueError("O valor deve ser positivo e razoável.")
    return round(amount,2)


@app.route("/eventos/<int:event_id>/orcamento", methods=["GET","POST"])
@login_required
def event_budget(event_id):
    event=get_event_or_404(event_id); db=get_db(); allowed={"previsto","em_cotacao","aprovado","contratado","realizado","cancelado"}
    if request.method=="POST":
        action=request.form.get("action","add"); item_id=request.form.get("item_id",type=int)
        if action=="delete":
            db.execute("DELETE FROM event_budget_items WHERE id=? AND event_id=?",(item_id,event_id)); db.commit(); flash("Item de orçamento removido.","success"); return redirect(url_for("event_budget",event_id=event_id))
        name=(request.form.get("item_name") or "").strip()
        if not name: flash("Informe o nome da despesa ou serviço.","error"); return redirect(url_for("event_budget",event_id=event_id))
        try: est=_parse_money(request.form.get("estimated_amount")); act=_parse_money(request.form.get("actual_amount"))
        except ValueError as exc: flash(str(exc),"error"); return redirect(url_for("event_budget",event_id=event_id))
        status=request.form.get("status","previsto"); status=status if status in allowed else "previsto"
        category=(request.form.get("category") or "Outros").strip()[:100] or "Outros"; supplier=(request.form.get("supplier") or "").strip()[:180] or None; notes=(request.form.get("notes") or "").strip()[:1500] or None; now=now_iso()
        vals=(name[:240],category,supplier,est,act,status,notes,now)
        if action=="edit" and item_id:
            db.execute("UPDATE event_budget_items SET item_name=?,category=?,supplier=?,estimated_amount=?,actual_amount=?,status=?,notes=?,updated_at=? WHERE id=? AND event_id=?",vals+(item_id,event_id)); flash("Item do orçamento atualizado.","success")
        else:
            db.execute("INSERT INTO event_budget_items(event_id,item_name,category,supplier,estimated_amount,actual_amount,status,notes,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?,?,?)",(event_id,name[:240],category,supplier,est,act,status,notes,now,now)); flash("Item incluído no orçamento do evento.","success")
        db.commit(); return redirect(url_for("event_budget",event_id=event_id))
    items=db.execute("SELECT * FROM event_budget_items WHERE event_id=? ORDER BY category,item_name,id",(event_id,)).fetchall()
    est=sum(float(x["estimated_amount"] or 0) for x in items); act=sum(float(x["actual_amount"] or 0) for x in items)
    return render_template("event_budget.html",event=event,items=items,total_estimated=est,total_actual=act,variance=act-est)


@app.route("/eventos/<int:event_id>/ocorrencias", methods=["GET","POST"])
@login_required
def event_incidents(event_id):
    event=get_event_or_404(event_id); db=get_db()
    categories=["Infraestrutura","Tecnologia / audiovisual","Fornecedor","Programação","Público / credenciamento","Comunicação","Acessibilidade / segurança","Outro"]
    statuses={"aberta":"Aberta","em_tratamento":"Em tratamento","resolvida":"Resolvida"}
    if request.method=="POST":
        action=request.form.get("action","add"); incident_id=request.form.get("incident_id",type=int)
        if action=="delete":
            db.execute("DELETE FROM event_incidents WHERE id=? AND event_id=?",(incident_id,event_id)); db.commit(); flash("Ocorrência removida do evento.","success"); return redirect(url_for("event_incidents",event_id=event_id))
        title=(request.form.get("title") or "").strip()
        if not title: flash("Informe um título curto para a ocorrência.","error"); return redirect(url_for("event_incidents",event_id=event_id))
        when=(request.form.get("incident_date") or "").strip() or None
        if when and not re.fullmatch(r"\d{4}-\d{2}-\d{2}(?:T\d{2}:\d{2})?",when): flash("Informe uma data/hora válida.","error"); return redirect(url_for("event_incidents",event_id=event_id))
        category=request.form.get("category","Outro"); category=category if category in categories else "Outro"; status=request.form.get("status","aberta"); status=status if status in statuses else "aberta"
        impact=(request.form.get("impact") or "").strip()[:2000] or None; resolution=(request.form.get("resolution") or "").strip()[:2000] or None; preventive=(request.form.get("preventive_action") or "").strip()[:2000] or None; responsible=(request.form.get("responsible") or "").strip()[:180] or None; now=now_iso(); user=session.get("user_email")
        if action=="edit" and incident_id:
            db.execute("UPDATE event_incidents SET incident_date=?,category=?,title=?,impact=?,resolution=?,preventive_action=?,responsible=?,status=?,recorded_by=?,updated_at=? WHERE id=? AND event_id=?",(when,category,title[:240],impact,resolution,preventive,responsible,status,user,now,incident_id,event_id)); flash("Ocorrência atualizada.","success")
        else:
            db.execute("INSERT INTO event_incidents(event_id,incident_date,category,title,impact,resolution,preventive_action,responsible,status,recorded_by,created_at,updated_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",(event_id,when,category,title[:240],impact,resolution,preventive,responsible,status,user,now,now)); flash("Ocorrência registrada no evento.","success")
        db.commit(); return redirect(url_for("event_incidents",event_id=event_id))
    incidents=db.execute("SELECT * FROM event_incidents WHERE event_id=? ORDER BY CASE status WHEN 'aberta' THEN 1 WHEN 'em_tratamento' THEN 2 ELSE 3 END, incident_date DESC,id DESC",(event_id,)).fetchall()
    counts={key:sum(1 for item in incidents if item["status"]==key) for key in statuses}
    return render_template("event_incidents.html",event=event,incidents=incidents,categories=categories,incident_statuses=statuses,counts=counts)


@app.get("/")
@login_required
def dashboard():
    db = get_db()
    counts = {r["status"]: r["total"] for r in db.execute("SELECT status, COUNT(*) total FROM events GROUP BY status").fetchall()}
    total = sum(counts.values())
    pending = counts.get("recebido", 0) + counts.get("em_triagem", 0) + counts.get("aguardando_informacoes", 0)
    active = counts.get("aprovado", 0) + counts.get("em_planejamento", 0)
    upcoming = db.execute(
        "SELECT * FROM events WHERE start_date >= ? AND status NOT IN ('cancelado', 'realizado', 'encerrado') ORDER BY start_date LIMIT 6",
        (date.today().isoformat(),),
    ).fetchall()
    recent = db.execute("SELECT * FROM events ORDER BY created_at DESC LIMIT 7").fetchall()
    future_events = db.execute("SELECT * FROM events WHERE status NOT IN ('cancelado','realizado','encerrado') AND start_date >= ? ORDER BY start_date LIMIT 100", (date.today().isoformat(),)).fetchall()
    schedule_conflicts = _find_schedule_conflicts(future_events)
    critical_tasks = _critical_task_rows(8)
    pending_confirmations = _pending_department_confirmations(8)
    alert_count = len(schedule_conflicts) + len(critical_tasks) + len(pending_confirmations)
    return render_template("dashboard.html", counts=counts, total=total, pending=pending, active=active, upcoming=upcoming, recent=recent,
                           schedule_conflicts=schedule_conflicts, critical_tasks=critical_tasks,
                           pending_confirmations=pending_confirmations, alert_count=alert_count)


def _resolve_internal_owner(db, event, form):
    """Resolve an internal owner only to an active registered account.

    Existing free-text owner values are preserved unless the select was changed,
    so updating status/notes does not silently erase legacy assignments. New
    assignments must always point to an active user account.
    """
    selected = (form.get("internal_owner") or "").strip()
    owner_changed = form.get("owner_changed") == "1"

    if selected == "__none__":
        return None, None
    if selected:
        user = db.execute(
            "SELECT email FROM users WHERE lower(email) = lower(?) AND is_active = 1",
            (selected,),
        ).fetchone()
        if not user:
            return None, "O responsável interno deve ser uma conta ativa cadastrada no sistema."
        return user["email"], None
    if owner_changed:
        return None, "Selecione uma conta ativa ou escolha 'Sem responsável'."
    return event["internal_owner"], None


def _active_internal_users(db):
    """Accounts that may be assigned as an internal event owner."""
    return db.execute(
        "SELECT id, email, role FROM users WHERE is_active = 1 ORDER BY email COLLATE NOCASE"
    ).fetchall()


@app.route("/triagem", methods=["GET", "POST"])
@login_required
def triage_queue():
    """Dedicated landing page for staff triage and internal follow-up."""
    db = get_db()
    if request.method == "POST":
        event_id = request.form.get("event_id", type=int)
        event = db.execute("SELECT * FROM events WHERE id = ?", (event_id,)).fetchone()
        if not event:
            abort(404)
        status = request.form.get("status", "")
        if status not in STATUS_LABELS:
            flash("Status inválido.", "error")
            return redirect(url_for("triage_queue"))
        if status == "encerrado":
            closure = db.execute("SELECT id FROM event_closures WHERE event_id = ?", (event_id,)).fetchone()
            if not closure:
                flash("Para encerrar um evento, registre primeiro o relatório de pós-evento.", "error")
                return redirect(url_for("finalize_event", event_id=event_id))
        if db.execute("SELECT id FROM event_closures WHERE event_id = ?", (event_id,)).fetchone() and status != "encerrado":
            flash("Este evento já tem relatório de encerramento. O status final foi preservado.", "error")
            return redirect(url_for("event_detail", event_id=event_id) + "#triagem")
        priority = request.form.get("priority", "Normal")
        if priority not in {"Baixa", "Normal", "Alta", "Urgente"}:
            priority = "Normal"
        owner, owner_error = _resolve_internal_owner(db, event, request.form)
        if owner_error:
            flash(owner_error, "error")
            return redirect(url_for("triage_queue", status=request.args.get("status", ""), q=request.args.get("q", "")))
        owner = owner[:180] if owner else None
        notes = (request.form.get("internal_notes") or "").strip()[:5000] or None
        db.execute(
            "UPDATE events SET status=?, priority=?, internal_owner=?, internal_notes=?, updated_at=? WHERE id=?",
            (status, priority, owner, notes, now_iso(), event_id),
        )
        db.commit()
        if status in {"aprovado", "em_planejamento"}:
            created_count = add_checklist(get_event_or_404(event_id))
            flash("Triagem salva. " + (f"Checklist inicial criado com {created_count} tarefas." if created_count else "O checklist existente foi mantido."), "success")
        else:
            flash("Triagem e responsável interno atualizados.", "success")
        return redirect(url_for("triage_queue", status=request.args.get("status", ""), q=request.args.get("q", "")))

    status_filter = (request.args.get("status") or "pendentes").strip()
    search = (request.args.get("q") or "").strip()[:160]
    sql = "SELECT * FROM events WHERE 1=1"
    params = []
    if status_filter == "pendentes":
        sql += " AND status IN ('recebido','em_triagem','aguardando_informacoes')"
    elif status_filter and status_filter != "todos" and status_filter in STATUS_LABELS:
        sql += " AND status = ?"
        params.append(status_filter)
    if search:
        like = f"%{search}%"
        sql += " AND (title LIKE ? OR protocol LIKE ? OR department LIKE ? OR organizer_name LIKE ?)"
        params.extend([like, like, like, like])
    sql += " ORDER BY CASE status WHEN 'recebido' THEN 1 WHEN 'em_triagem' THEN 2 WHEN 'aguardando_informacoes' THEN 3 WHEN 'aprovado' THEN 4 WHEN 'em_planejamento' THEN 5 ELSE 6 END, start_date ASC, id DESC"
    events = db.execute(sql, params).fetchall()
    counts = {row["status"]: row["total"] for row in db.execute("SELECT status, COUNT(*) total FROM events GROUP BY status").fetchall()}
    active_users = _active_internal_users(db)
    return render_template("triage.html", events=events, selected_status=status_filter, search=search, counts=counts, active_users=active_users)


@app.get("/eventos")
@login_required
def events_list():
    status = request.args.get("status", "")
    search = request.args.get("q", "").strip()
    sql = "SELECT * FROM events WHERE 1=1"
    params = []
    if status in STATUS_LABELS:
        sql += " AND status = ?"
        params.append(status)
    if search:
        like = f"%{search}%"
        sql += " AND (title LIKE ? OR protocol LIKE ? OR department LIKE ? OR organizer_name LIKE ?)"
        params.extend([like, like, like, like])
    sql += " ORDER BY CASE status WHEN 'recebido' THEN 1 WHEN 'em_triagem' THEN 2 WHEN 'aguardando_informacoes' THEN 3 WHEN 'aprovado' THEN 4 WHEN 'em_planejamento' THEN 5 ELSE 6 END, start_date ASC"
    events = get_db().execute(sql, params).fetchall()
    return render_template("events.html", events=events, selected_status=status, search=search)


@app.get("/eventos/<int:event_id>")
@login_required
def event_detail(event_id):
    event = get_event_or_404(event_id)
    tasks = get_db().execute("SELECT * FROM tasks WHERE event_id = ? ORDER BY sort_order, id", (event_id,)).fetchall()
    completed = sum(1 for task in tasks if task["status"] in TASK_DONE_STATUSES)
    progress = round((completed / len(tasks)) * 100) if tasks else 0
    db = get_db()
    closure = db.execute("SELECT * FROM event_closures WHERE event_id = ?", (event_id,)).fetchone()
    evaluation = db.execute("SELECT * FROM event_evaluations WHERE event_id = ?", (event_id,)).fetchone()
    date_from = (date.fromisoformat(event["start_date"]) - timedelta(days=45)).isoformat()
    date_to = (date.fromisoformat(event["end_date"] or event["start_date"]) + timedelta(days=45)).isoformat()
    conflict_pool = db.execute("SELECT * FROM events WHERE status<>'cancelado' AND start_date<=? AND COALESCE(NULLIF(end_date,''),start_date)>=? ORDER BY start_date,start_time", (date_to,date_from)).fetchall()
    schedule_conflicts = [c for c in _find_schedule_conflicts(conflict_pool) if event_id in {c["event_a_id"],c["event_b_id"]}]
    program_count = db.execute("SELECT COUNT(*) FROM event_program_items WHERE event_id=?",(event_id,)).fetchone()[0]
    incident_count = db.execute("SELECT COUNT(*) FROM event_incidents WHERE event_id=?",(event_id,)).fetchone()[0]
    open_incident_count = db.execute("SELECT COUNT(*) FROM event_incidents WHERE event_id=? AND status<>'resolvida'",(event_id,)).fetchone()[0]
    budget = db.execute("SELECT COALESCE(SUM(estimated_amount),0) estimated,COALESCE(SUM(actual_amount),0) actual FROM event_budget_items WHERE event_id=?",(event_id,)).fetchone()
    return render_template(
        "event_detail.html", event=event, tasks=tasks, completed=completed, progress=progress,
        event_support=event_support(event), event_modes=event_modes(event),
        event_venues=event_venues(event), event_items=event_requested_items(event),
        task_statuses=TASK_STATUSES, closure=closure, evaluation=evaluation,
        schedule_conflicts=schedule_conflicts, program_count=program_count, incident_count=incident_count,
        open_incident_count=open_incident_count, budget_estimated=budget["estimated"], budget_actual=budget["actual"],
        active_users=_active_internal_users(db),
        available_task_templates=db.execute("SELECT * FROM task_templates WHERE active=1 ORDER BY category, sort_order, title").fetchall(),
        existing_task_keys={(str(t["title"]).strip().casefold(), str(t["category"] or "").strip().casefold()) for t in tasks},
    )


def _report_filters(args):
    raw_year = (args.get("ano") or str(date.today().year)).strip()
    try:
        year = int(raw_year)
        if year < 2000 or year > 2100:
            raise ValueError
    except ValueError:
        year = date.today().year
    status_filter = (args.get("status") or "realizados").strip()
    allowed_filters = {"realizados", "encerrados", "todos", *STATUS_LABELS.keys()}
    if status_filter not in allowed_filters:
        status_filter = "realizados"
    event_type = (args.get("categoria") or "").strip()
    department = (args.get("setor") or "").strip()
    return year, status_filter, event_type, department


def _annual_report_rows(year: int, status_filter: str, event_type: str = "", department: str = ""):
    db = get_db()
    start = f"{year:04d}-01-01"
    end = f"{year + 1:04d}-01-01"
    sql = """SELECT e.*, c.actual_attendees, c.summary AS closure_summary,
        c.positives AS closure_positives, c.issues AS closure_issues,
        c.improvements AS closure_improvements, c.follow_up_actions AS closure_follow_up,
        c.final_observations AS closure_observations, c.closed_by, c.closed_at,
        ev.rating_overall, ev.rating_sector, ev.rating_communication,
        ev.rating_responsiveness, ev.rating_needs_met, ev.positive_feedback,
        ev.improvement_suggestions, ev.additional_comments, ev.responded_at,
        (SELECT COALESCE(SUM(bi.estimated_amount),0) FROM event_budget_items bi WHERE bi.event_id=e.id) AS budget_estimated_total,
        (SELECT COALESCE(SUM(bi.actual_amount),0) FROM event_budget_items bi WHERE bi.event_id=e.id) AS budget_actual_total,
        (SELECT COUNT(*) FROM event_budget_items bi WHERE bi.event_id=e.id AND bi.estimated_amount IS NOT NULL) AS budget_estimated_count,
        (SELECT COUNT(*) FROM event_budget_items bi WHERE bi.event_id=e.id AND bi.actual_amount IS NOT NULL) AS budget_actual_count,
        (SELECT COUNT(*) FROM event_incidents ei WHERE ei.event_id=e.id) AS incident_count,
        (SELECT COUNT(*) FROM event_incidents ei WHERE ei.event_id=e.id AND ei.status<>'resolvida') AS open_incident_count
        FROM events e
        LEFT JOIN event_closures c ON c.event_id=e.id
        LEFT JOIN event_evaluations ev ON ev.event_id=e.id
        WHERE e.start_date >= ? AND e.start_date < ?"""
    params = [start, end]
    if status_filter == "realizados":
        sql += " AND e.status IN ('realizado','encerrado')"
    elif status_filter == "encerrados":
        sql += " AND e.status = 'encerrado'"
    elif status_filter != "todos":
        sql += " AND e.status = ?"
        params.append(status_filter)
    if event_type:
        if event_type == "Outro":
            sql += " AND e.event_type = ?"
            params.append(event_type)
        else:
            sql += " AND (e.event_type = ? OR e.event_type_other = ?)"
            params.extend([event_type, event_type])
    if department:
        sql += " AND e.department = ?"
        params.append(department)
    sql += " ORDER BY e.start_date, e.start_time, e.title"
    return db.execute(sql, params).fetchall()


def _rating_average(rows, field):
    values = [int(row[field]) for row in rows if row[field] is not None]
    return round(sum(values) / len(values), 1) if values else None


@app.get("/relatorios")
@login_required
def event_reports():
    year, status_filter, event_type, department = _report_filters(request.args)
    rows = _annual_report_rows(year, status_filter, event_type, department)
    db = get_db()
    departments = [r["department"] for r in db.execute(
        "SELECT DISTINCT department FROM events WHERE department IS NOT NULL AND department <> '' ORDER BY department"
    ).fetchall()]
    finished = [r for r in rows if r["status"] in {"realizado", "encerrado"}]
    closed = [r for r in rows if r["status"] == "encerrado"]
    actual_attendance = sum(r["actual_attendees"] or 0 for r in rows)
    evaluated = sum(1 for r in rows if r["responded_at"])
    budget_estimated_total = sum(float(r["budget_estimated_total"] or 0) for r in rows)
    budget_actual_total = sum(float(r["budget_actual_total"] or 0) for r in rows)
    incident_total = sum(int(r["incident_count"] or 0) for r in rows)
    incident_open_total = sum(int(r["open_incident_count"] or 0) for r in rows)
    budget_estimated_count = sum(int(r["budget_estimated_count"] or 0) for r in rows)
    budget_actual_count = sum(int(r["budget_actual_count"] or 0) for r in rows)
    return render_template(
        "reports.html", rows=rows, year=year, status_filter=status_filter,
        event_type_filter=event_type, department_filter=department,
        departments=departments, event_types=EVENT_TYPES,
        finished_count=len(finished), closed_count=len(closed),
        actual_attendance=actual_attendance, evaluated_count=evaluated,
        budget_estimated_total=budget_estimated_total, budget_actual_total=budget_actual_total,
        budget_estimated_count=budget_estimated_count, budget_actual_count=budget_actual_count,
        incident_total=incident_total, incident_open_total=incident_open_total,
        avg_overall=_rating_average(rows, "rating_overall"),
        avg_sector=_rating_average(rows, "rating_sector"),
    )



@app.get("/relatorios/comparativo-custos")
@login_required
def budget_comparison_report():
    """Comparativo histórico dos valores previstos, contratados e realizados."""
    db = get_db()
    selected_type = (request.args.get("event_type") or "").strip()
    if selected_type and selected_type not in EVENT_TYPES:
        # Permite também categorias livres registradas em event_type_other.
        valid_other = db.execute(
            "SELECT 1 FROM events WHERE event_type_other = ? LIMIT 1", (selected_type,)
        ).fetchone()
        if not valid_other:
            selected_type = ""

    where = []
    params = []
    if selected_type:
        if selected_type == "Outro":
            where.append("e.event_type = ?")
            params.append("Outro")
        else:
            where.append("(e.event_type = ? OR e.event_type_other = ?)")
            params.extend([selected_type, selected_type])
    where_sql = (" WHERE " + " AND ".join(where)) if where else ""
    rows = db.execute(f"""
        SELECT e.id, e.protocol, e.title, e.start_date, e.event_type, e.event_type_other,
               COALESCE((SELECT SUM(bi.estimated_amount) FROM event_budget_items bi
                         WHERE bi.event_id = e.id), 0) AS estimated_total,
               COALESCE((SELECT SUM(CASE WHEN bi.status IN ('contratado','realizado')
                                          THEN COALESCE(bi.estimated_amount, 0) ELSE 0 END)
                         FROM event_budget_items bi WHERE bi.event_id = e.id), 0) AS contracted_total,
               COALESCE((SELECT SUM(bi.actual_amount) FROM event_budget_items bi
                         WHERE bi.event_id = e.id), 0) AS actual_total,
               (SELECT COUNT(*) FROM event_budget_items bi WHERE bi.event_id = e.id) AS expense_count
        FROM events e
        {where_sql}
        ORDER BY e.start_date DESC, e.title COLLATE NOCASE
    """, params).fetchall()

    grouped_map = {}
    for row in rows:
        year = (row["start_date"] or "")[:4] or "Sem data"
        kind = (row["event_type_other"] or row["event_type"] or "Não informado")
        key = (year, kind)
        if key not in grouped_map:
            grouped_map[key] = {"year": year, "kind": kind, "count": 0,
                                "estimated": 0.0, "contracted": 0.0, "actual": 0.0}
        group = grouped_map[key]
        group["count"] += 1
        group["estimated"] += float(row["estimated_total"] or 0)
        group["contracted"] += float(row["contracted_total"] or 0)
        group["actual"] += float(row["actual_total"] or 0)

    grouped = sorted(grouped_map.values(),
                     key=lambda item: (item["year"], item["kind"]), reverse=True)
    return render_template(
        "budget_comparison.html", rows=rows, grouped=grouped,
        event_types=EVENT_TYPES, selected_type=selected_type,
    )

def _safe_csv_cell(value):
    """Avoid spreadsheet formula execution for exported user-entered text."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


@app.get("/relatorios/exportar.csv")
@login_required
def export_annual_report():
    year, status_filter, event_type, department = _report_filters(request.args)
    rows = _annual_report_rows(year, status_filter, event_type, department)
    output = io.StringIO(newline="")
    output.write("\ufeff")  # BOM makes Portuguese text display correctly in Excel.
    writer = csv.writer(output, delimiter=";")
    writer.writerow([
        "Protocolo", "Evento", "Data", "Horário de início", "Horário de término", "Categoria", "Setor solicitante", "Responsável",
        "Formato", "Local", "Público previsto", "Público real", "Status",
        "Resumo pós-evento", "Pontos positivos", "Falhas/dificuldades", "Melhorias sugeridas",
        "Ações de acompanhamento", "Avaliação geral (1-5)", "Avaliação Setor de Eventos (1-5)",
        "Comunicação (1-5)", "Agilidade (1-5)", "Necessidades atendidas (1-5)",
        "Pontos positivos do solicitante", "Sugestões do solicitante", "Comentários", "Respondida em", "Encerrado em",
        "Orçamento previsto (R$)", "Custo realizado (R$)", "Ocorrências registradas", "Ocorrências em aberto",
    ])
    for row in rows:
        csv_values = [
            row["protocol"], row["title"], row["start_date"], row["start_time"], row["end_time"],
            row["event_type_other"] or row["event_type"], row["department"], row["organizer_name"], row["event_format"], row["venue"],
            row["expected_attendees"], row["actual_attendees"], STATUS_LABELS.get(row["status"], row["status"]),
            row["closure_summary"], row["closure_positives"], row["closure_issues"], row["closure_improvements"],
            row["closure_follow_up"], row["rating_overall"], row["rating_sector"],
            row["rating_communication"], row["rating_responsiveness"], row["rating_needs_met"],
            row["positive_feedback"], row["improvement_suggestions"], row["additional_comments"],
            row["responded_at"], row["closed_at"], row["budget_estimated_total"] if row["budget_estimated_count"] else "",
            row["budget_actual_total"] if row["budget_actual_count"] else "",
            row["incident_count"], row["open_incident_count"],
        ]
        writer.writerow([_safe_csv_cell(value) for value in csv_values])
    filename = f"relatorio_eventos_fmabc_{year}.csv"
    return Response(
        output.getvalue(), content_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


def _public_app_base_url():
    configured = os.getenv("APP_BASE_URL", "").strip().rstrip("/")
    if configured:
        parsed = urlparse(configured)
        if parsed.scheme not in {"http", "https"} or not parsed.netloc or parsed.hostname in {"127.0.0.1", "localhost", "0.0.0.0"}:
            return ""
        return configured
    host = (request.host or "").split(":", 1)[0].lower()
    if host in {"127.0.0.1", "localhost", "0.0.0.0"}:
        return ""
    return request.url_root.rstrip("/")


def _send_evaluation_invitation(event, token: str):
    cfg = smtp_settings()
    if not smtp_is_configured():
        return False, "O SMTP ainda não está configurado. O link foi criado e pode ser copiado para envio manual."
    base_url = _public_app_base_url()
    if not base_url:
        return False, "Como o sistema está em endereço local, o link não é acessível ao solicitante. Configure APP_BASE_URL após publicar o sistema para enviar o convite por e-mail."
    recipient = (event["organizer_email"] or "").strip()
    if not recipient or "@" not in recipient:
        return False, "O e-mail do solicitante não está válido; copie o link para envio manual."
    evaluation_url = base_url + url_for("public_event_evaluation", token=token)
    msg = EmailMessage()
    msg["Subject"] = f"Avaliação do evento: {event['title']} — FMABC"
    msg["From"] = cfg["from_email"]
    msg["To"] = recipient
    msg.set_content(
        f"Olá, {event['organizer_name']}.\n\n"
        f"O evento {event['title']} ({event['protocol']}) foi encerrado pelo Setor de Eventos da FMABC. "
        "Sua avaliação é importante para aprimorar nosso atendimento e a organização de futuros eventos.\n\n"
        f"Responda à avaliação neste link: {evaluation_url}\n\n"
        "O formulário leva poucos minutos. Obrigado pela colaboração.\n\nSetor de Eventos — Centro Universitário FMABC"
    )
    try:
        with smtplib.SMTP(cfg["host"], cfg["port"], timeout=20) as server:
            if cfg["use_tls"]:
                server.starttls()
            if cfg["user"]:
                server.login(cfg["user"], cfg["password"])
            server.send_message(msg)
        return True, f"Convite de avaliação enviado para {recipient}."
    except Exception as exc:
        return False, f"Falha no envio SMTP ({type(exc).__name__}). O link continua disponível para envio manual."


@app.route("/eventos/<int:event_id>/finalizar", methods=["GET", "POST"])
@login_required
def finalize_event(event_id):
    db = get_db()
    event = get_event_or_404(event_id)
    if event["status"] == "cancelado":
        flash("Um evento cancelado não pode ser finalizado como realizado.", "error")
        return redirect(url_for("event_detail", event_id=event_id))
    if request.method == "POST":
        if request.form.get("confirm_realizado") != "yes":
            flash("Confirme que o evento aconteceu antes de registrar o encerramento.", "error")
            return redirect(url_for("finalize_event", event_id=event_id))
        summary = request.form.get("summary", "").strip()
        if not summary:
            flash("Preencha o resumo do que aconteceu no evento.", "error")
            return redirect(url_for("finalize_event", event_id=event_id))
        raw_attendees = request.form.get("actual_attendees", "").strip()
        actual_attendees = None
        if raw_attendees:
            try:
                actual_attendees = int(raw_attendees)
                if actual_attendees < 0:
                    raise ValueError
            except ValueError:
                flash("Informe um número válido para o público real ou deixe o campo em branco.", "error")
                return redirect(url_for("finalize_event", event_id=event_id))
        now = now_iso()
        existing_closure = db.execute("SELECT id FROM event_closures WHERE event_id=?", (event_id,)).fetchone()
        if existing_closure:
            db.execute(
                """UPDATE event_closures SET actual_attendees=?, summary=?, positives=?, issues=?, improvements=?,
                   follow_up_actions=?, final_observations=?, updated_by=?, updated_at=? WHERE event_id=?""",
                (actual_attendees, summary, request.form.get("positives", "").strip() or None,
                 request.form.get("issues", "").strip() or None, request.form.get("improvements", "").strip() or None,
                 request.form.get("follow_up_actions", "").strip() or None,
                 request.form.get("final_observations", "").strip() or None,
                 session.get("user_email"), now, event_id),
            )
        else:
            db.execute(
                """INSERT INTO event_closures(event_id, actual_attendees, summary, positives, issues, improvements,
                   follow_up_actions, final_observations, closed_by, closed_at, updated_by, updated_at)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?,?)""",
                (event_id, actual_attendees, summary, request.form.get("positives", "").strip() or None,
                 request.form.get("issues", "").strip() or None, request.form.get("improvements", "").strip() or None,
                 request.form.get("follow_up_actions", "").strip() or None,
                 request.form.get("final_observations", "").strip() or None,
                 session.get("user_email"), now, session.get("user_email"), now),
            )
        db.execute("UPDATE events SET status='encerrado', updated_at=? WHERE id=?", (now, event_id))
        evaluation = db.execute("SELECT * FROM event_evaluations WHERE event_id=?", (event_id,)).fetchone()
        if not evaluation:
            token = secrets.token_urlsafe(32)
            db.execute(
                "INSERT INTO event_evaluations(event_id, token, requester_email, created_at) VALUES (?,?,?,?)",
                (event_id, token, event["organizer_email"], now),
            )
        db.commit()
        evaluation = db.execute("SELECT * FROM event_evaluations WHERE event_id=?", (event_id,)).fetchone()
        if request.form.get("send_evaluation") == "yes" and not evaluation["responded_at"]:
            refreshed_event = get_event_or_404(event_id)
            sent, message = _send_evaluation_invitation(refreshed_event, evaluation["token"])
            if sent:
                db.execute("UPDATE event_evaluations SET invitation_sent_at=?, invitation_method='smtp' WHERE event_id=?", (now_iso(), event_id))
                db.commit()
                flash("Encerramento registrado. " + message, "success")
            else:
                flash("Encerramento registrado. " + message, "error")
        else:
            flash("Encerramento registrado. O link para avaliação está disponível nesta página.", "success")
        return redirect(url_for("finalize_event", event_id=event_id))
    closure = db.execute("SELECT * FROM event_closures WHERE event_id=?", (event_id,)).fetchone()
    evaluation = db.execute("SELECT * FROM event_evaluations WHERE event_id=?", (event_id,)).fetchone()
    return render_template(
        "event_closure.html", event=event, closure=closure, evaluation=evaluation,
        evaluation_url=(url_for("public_event_evaluation", token=evaluation["token"], _external=True) if evaluation else None),
        smtp_configured=smtp_is_configured(), public_base_configured=bool(os.getenv("APP_BASE_URL", "").strip()),
    )


@app.route("/avaliacao-evento/<token>", methods=["GET", "POST"])
def public_event_evaluation(token):
    db = get_db()
    evaluation = db.execute(
        "SELECT ev.*, e.title AS event_title, e.protocol, e.start_date FROM event_evaluations ev JOIN events e ON e.id=ev.event_id WHERE ev.token=?",
        (token,),
    ).fetchone()
    if not evaluation:
        abort(404)
    submitted = False
    if request.method == "POST":
        if evaluation["responded_at"]:
            submitted = True
        else:
            rating_fields = ["rating_overall", "rating_sector", "rating_communication", "rating_responsiveness", "rating_needs_met"]
            ratings = {}
            try:
                for field in rating_fields:
                    value = int(request.form.get(field, ""))
                    if value not in {1, 2, 3, 4, 5}:
                        raise ValueError
                    ratings[field] = value
            except ValueError:
                flash("Selecione uma nota de 1 a 5 em todas as perguntas de avaliação.", "error")
                return render_template("evaluation_public.html", evaluation=evaluation, submitted=False), 400
            db.execute(
                """UPDATE event_evaluations SET rating_overall=?, rating_sector=?, rating_communication=?,
                   rating_responsiveness=?, rating_needs_met=?, positive_feedback=?, improvement_suggestions=?,
                   additional_comments=?, responded_at=? WHERE id=? AND responded_at IS NULL""",
                (ratings["rating_overall"], ratings["rating_sector"], ratings["rating_communication"],
                 ratings["rating_responsiveness"], ratings["rating_needs_met"],
                 request.form.get("positive_feedback", "").strip() or None,
                 request.form.get("improvement_suggestions", "").strip() or None,
                 request.form.get("additional_comments", "").strip() or None,
                 now_iso(), evaluation["id"]),
            )
            db.commit()
            submitted = True
    return render_template("evaluation_public.html", evaluation=evaluation, submitted=submitted)


@app.get("/eventos/<int:event_id>/checklist/imprimir")
@login_required
def print_event_checklist(event_id):
    """Render a compact checklist designed to print on no more than one A4 sheet, duplex."""
    event = get_event_or_404(event_id)
    tasks = list(get_db().execute(
        "SELECT * FROM tasks WHERE event_id = ? ORDER BY sort_order, id", (event_id,)
    ).fetchall())
    if not tasks:
        flash("Este evento ainda não possui checklist para imprimir.", "error")
        return redirect(url_for("event_detail", event_id=event_id) + "#checklist")

    # Short checklists fit on one side. Longer checklists are explicitly divided into
    # two printable sheets so that duplex printing uses at most one A4 physical sheet.
    if len(tasks) <= 24:
        task_pages = [tasks]
    else:
        midpoint = (len(tasks) + 1) // 2
        task_pages = [tasks[:midpoint], tasks[midpoint:]]
    completed = sum(1 for task in tasks if task["status"] in TASK_DONE_STATUSES)
    progress = round((completed / len(tasks)) * 100)
    max_page_tasks = max(len(page) for page in task_pages)
    density = "ultra" if max_page_tasks > 28 else ("dense" if max_page_tasks > 22 else "normal")
    return render_template(
        "checklist_print.html", event=event, tasks=tasks, task_pages=task_pages,
        task_statuses=TASK_STATUSES, completed=completed, progress=progress,
        auto_print=request.args.get("auto") == "1", density=density,
    )


@app.post("/eventos/<int:event_id>/status")
@login_required
def update_event_status(event_id):
    event = get_event_or_404(event_id)
    status = request.form.get("status", "")
    if status not in STATUS_LABELS:
        flash("Status inválido.", "error")
        return redirect(url_for("event_detail", event_id=event_id))
    existing_closure = get_db().execute("SELECT id FROM event_closures WHERE event_id=?", (event_id,)).fetchone()
    if status == "encerrado" and not existing_closure:
        flash("Para encerrar um evento, registre primeiro o relatório de pós-evento.", "error")
        return redirect(url_for("finalize_event", event_id=event_id))
    if existing_closure and status != "encerrado":
        flash("Este evento já tem um relatório de encerramento. Atualize o relatório de pós-evento para corrigir informações; o status final foi preservado.", "error")
        return redirect(url_for("event_detail", event_id=event_id))
    db = get_db()
    owner, owner_error = _resolve_internal_owner(db, event, request.form)
    if owner_error:
        flash(owner_error, "error")
        return redirect(url_for("event_detail", event_id=event_id) + "#triagem")
    owner = owner[:180] if owner else None
    db.execute(
        "UPDATE events SET status = ?, priority = ?, internal_owner = ?, internal_notes = ?, updated_at = ? WHERE id = ?",
        (status, request.form.get("priority", "Normal") if request.form.get("priority") in {"Baixa", "Normal", "Alta", "Urgente"} else "Normal",
         owner, request.form.get("internal_notes", "").strip() or None, now_iso(), event_id),
    )
    db.commit()
    if status in {"aprovado", "em_planejamento"}:
        refreshed = get_event_or_404(event_id)
        # Aviso preventivo: conflitos precisam de conferência humana, não bloqueiam o fluxo.
        try:
            overlaps = _find_schedule_conflicts(db.execute("SELECT * FROM events WHERE status<>'cancelado' AND start_date BETWEEN ? AND ? ORDER BY start_date,start_time", ((date.fromisoformat(event["start_date"])-timedelta(days=1)).isoformat(), (date.fromisoformat(event["end_date"] or event["start_date"])+timedelta(days=1)).isoformat())).fetchall())
            if any(event_id in {c["event_a_id"],c["event_b_id"]} for c in overlaps):
                flash("Atenção: há possível conflito de local/horário. Confira a Agenda Institucional.", "warning")
        except (ValueError, TypeError):
            pass
        created_count = add_checklist(refreshed)
        if created_count:
            flash(f"Status atualizado. Checklist inicial criado com {created_count} tarefas.", "success")
        else:
            flash("Status atualizado. O checklist existente foi mantido.", "success")
    else:
        flash("Status e informações internas atualizados.", "success")
    return redirect(url_for("event_detail", event_id=event_id))


@app.post("/eventos/<int:event_id>/tarefas/<int:task_id>/alternar")
@login_required
def toggle_task(event_id, task_id):
    get_event_or_404(event_id)
    db = get_db()
    task = db.execute("SELECT * FROM tasks WHERE id = ? AND event_id = ?", (task_id, event_id)).fetchone()
    if not task:
        abort(404)
    completed = task["status"] != "concluida"
    db.execute("UPDATE tasks SET status = ?, completed_at = ? WHERE id = ?",
               ("concluida" if completed else "pendente", now_iso() if completed else None, task_id))
    db.commit()
    return redirect(url_for("event_detail", event_id=event_id) + "#checklist")


@app.post("/eventos/<int:event_id>/tarefas/<int:task_id>/editar")
@login_required
def edit_task(event_id, task_id):
    get_event_or_404(event_id)
    db = get_db()
    task = db.execute("SELECT * FROM tasks WHERE id = ? AND event_id = ?", (task_id, event_id)).fetchone()
    if not task:
        abort(404)
    title = request.form.get("title", "").strip()
    status = request.form.get("status", "pendente")
    if not title:
        flash("O título da tarefa é obrigatório.", "error")
        return redirect(url_for("event_detail", event_id=event_id) + "#checklist")
    if status not in TASK_STATUSES:
        flash("Status da tarefa inválido.", "error")
        return redirect(url_for("event_detail", event_id=event_id) + "#checklist")
    due = normalize_date(request.form.get("due_date"))
    if (request.form.get("due_date") or "").strip() and not due:
        flash("Informe um prazo válido.", "error")
        return redirect(url_for("event_detail", event_id=event_id) + "#checklist")
    db.execute(
        """UPDATE tasks SET title = ?, category = ?, due_date = ?, status = ?, assigned_to = ?,
           description = ?, notes = ?, completed_at = ? WHERE id = ? AND event_id = ?""",
        (title[:240], request.form.get("category", "Organização").strip()[:80] or "Organização", due, status,
         request.form.get("assigned_to", "").strip()[:180] or None,
         request.form.get("description", "").strip()[:1500] or None,
         request.form.get("notes", "").strip()[:3000] or None,
         now_iso() if status == "concluida" else None, task_id, event_id),
    )
    db.execute("UPDATE events SET updated_at = ? WHERE id = ?", (now_iso(), event_id))
    db.commit()
    flash("Tarefa atualizada.", "success")
    return redirect(url_for("event_detail", event_id=event_id) + "#task-" + str(task_id))


@app.post("/eventos/<int:event_id>/tarefas/adicionar")
@login_required
def add_task(event_id):
    get_event_or_404(event_id)
    title = request.form.get("title", "").strip()
    if not title:
        flash("Digite o nome da tarefa.", "error")
        return redirect(url_for("event_detail", event_id=event_id) + "#checklist")
    due = normalize_date(request.form.get("due_date"))
    if (request.form.get("due_date") or "").strip() and not due:
        flash("Informe um prazo válido.", "error")
        return redirect(url_for("event_detail", event_id=event_id) + "#checklist")
    status = request.form.get("status", "pendente")
    if status not in TASK_STATUSES:
        status = "pendente"
    db = get_db()
    next_order = db.execute("SELECT COALESCE(MAX(sort_order), 0) + 1 n FROM tasks WHERE event_id = ?", (event_id,)).fetchone()["n"]
    db.execute(
        """INSERT INTO tasks(event_id, title, category, due_date, status, assigned_to, description, notes, sort_order, created_at, completed_at)
           VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
        (event_id, title[:240], request.form.get("category", "Organização").strip()[:80] or "Organização", due, status,
         request.form.get("assigned_to", "").strip()[:180] or None,
         request.form.get("description", "").strip()[:1500] or None,
         request.form.get("notes", "").strip()[:3000] or None, next_order, now_iso(), now_iso() if status == "concluida" else None),
    )
    db.execute("UPDATE events SET updated_at = ? WHERE id = ?", (now_iso(), event_id))
    db.commit()
    flash("Tarefa adicionada ao checklist.", "success")
    return redirect(url_for("event_detail", event_id=event_id) + "#checklist")


@app.post("/eventos/<int:event_id>/tarefas/adicionar-modelos")
@login_required
def add_tasks_from_library(event_id):
    """Add selected active library templates to this event without changing the shared model."""
    event = get_event_or_404(event_id)
    db = get_db()
    selected_ids = []
    for raw_id in request.form.getlist("template_ids"):
        try:
            selected_ids.append(int(raw_id))
        except (TypeError, ValueError):
            continue
    if not selected_ids:
        flash("Selecione pelo menos uma atividade da biblioteca.", "error")
        return redirect(url_for("event_detail", event_id=event_id) + "#checklist")

    event_day = date.fromisoformat(event["start_date"])
    existing = {(str(row["title"]).strip().casefold(), str(row["category"] or "").strip().casefold())
                for row in db.execute("SELECT title, category FROM tasks WHERE event_id=?", (event_id,)).fetchall()}
    rows = db.execute(
        f"SELECT * FROM task_templates WHERE active=1 AND id IN ({','.join('?' for _ in selected_ids)}) ORDER BY category, sort_order, title",
        selected_ids,
    ).fetchall()
    added = 0
    skipped = 0
    next_order = db.execute("SELECT COALESCE(MAX(sort_order),0)+1 AS n FROM tasks WHERE event_id=?", (event_id,)).fetchone()["n"]
    for template in rows:
        key = (str(template["title"]).strip().casefold(), str(template["category"] or "").strip().casefold())
        if key in existing:
            skipped += 1
            continue
        due_date = (event_day + timedelta(days=int(template["due_offset"] or 0))).isoformat()
        db.execute("""INSERT INTO tasks(event_id,title,category,due_date,status,assigned_to,description,notes,sort_order,created_at)
                      VALUES(?,?,?,?, 'pendente', ?, ?, NULL, ?, ?)""",
                   (event_id, template["title"], template["category"], due_date,
                    template["assigned_to"], template["description"], next_order, now_iso()))
        next_order += 1
        existing.add(key)
        added += 1
    db.execute("UPDATE events SET updated_at=? WHERE id=?", (now_iso(), event_id))
    db.commit()
    message = f"Adicionadas {added} atividade(s) da biblioteca ao checklist."
    if skipped:
        message += f" {skipped} já existiam e foram ignoradas para evitar duplicidade."
    flash(message, "success" if added else "warning")
    return redirect(url_for("event_detail", event_id=event_id) + "#checklist")



@app.get("/checklist-consolidado")
@login_required
def consolidated_checklist():
    """Cross-event checklist view for the Events team."""
    db = get_db()
    event_id = request.args.get("evento", type=int)
    status = request.args.get("status", "").strip()
    responsible = request.args.get("responsavel", "").strip()
    search = request.args.get("q", "").strip()
    period = request.args.get("prazo", "").strip()
    sql = "SELECT t.*, e.title AS event_title, e.protocol AS event_protocol, e.start_date AS event_date FROM tasks t JOIN events e ON e.id=t.event_id WHERE 1=1"
    params = []
    if event_id:
        sql += " AND e.id = ?"; params.append(event_id)
    if status in TASK_STATUSES:
        sql += " AND t.status = ?"; params.append(status)
    if responsible:
        sql += " AND COALESCE(t.assigned_to,'') = ?"; params.append(responsible)
    if search:
        like = f"%{search}%"; sql += " AND (t.title LIKE ? OR e.title LIKE ? OR e.protocol LIKE ?)"; params.extend([like, like, like])
    today = date.today().isoformat()
    if period == "atrasadas":
        sql += " AND t.due_date IS NOT NULL AND t.due_date<>'' AND t.due_date < ? AND t.status NOT IN ('concluida','nao_se_aplica')"; params.append(today)
    elif period == "7dias":
        end = (date.today() + timedelta(days=7)).isoformat(); sql += " AND t.due_date IS NOT NULL AND t.due_date<>'' AND t.due_date BETWEEN ? AND ? AND t.status NOT IN ('concluida','nao_se_aplica')"; params.extend([today, end])
    sql += " ORDER BY CASE WHEN t.due_date IS NULL OR t.due_date='' THEN 1 ELSE 0 END, t.due_date, e.start_date, t.sort_order, t.id"
    tasks = [dict(r) for r in db.execute(sql, params).fetchall()]
    for task in tasks:
        task['overdue'] = bool(task.get('due_date') and task['due_date'] < today and task['status'] not in TASK_DONE_STATUSES)
    events = db.execute("SELECT id,title,protocol FROM events WHERE status <> 'cancelado' ORDER BY start_date DESC").fetchall()
    people = db.execute("SELECT DISTINCT assigned_to FROM tasks WHERE assigned_to IS NOT NULL AND TRIM(assigned_to)<>'' ORDER BY assigned_to").fetchall()
    counts = {'total': len(tasks), 'pendentes': sum(1 for t in tasks if t['status'] not in TASK_DONE_STATUSES), 'atrasadas': sum(1 for t in tasks if t['overdue']), 'concluidas': sum(1 for t in tasks if t['status'] == 'concluida')}
    return render_template('checklist_consolidado.html', tasks=tasks, events=events, people=people, selected_event=event_id, selected_status=status, selected_responsible=responsible, selected_period=period, search=search, counts=counts, task_statuses=TASK_STATUSES)

@app.route("/modelos-checklist", methods=["GET", "POST"])
@login_required
def checklist_templates():
    if session.get("user_role") != "admin":
        abort(404)
    db = get_db()
    categories = ["Planejamento", "Administrativo", "Infraestrutura", "Manutenção", "Limpeza", "TI / Audiovisual", "Comunicação", "Inscrições", "Credenciamento", "Certificados", "Fornecedores", "Materiais", "Cerimonial", "Convidados", "Recepção", "Logística", "Transmissão", "Dia do evento", "Pós-evento", "Operação", "Outro"]
    trigger_types = {"always": "Todos os eventos", "support": "Serviço de apoio", "item": "Item solicitado", "mode": "Modalidade", "category": "Categoria do evento"}
    trigger_values = list(dict.fromkeys(SUPPORT_OPTIONS + REQUESTED_ITEM_OPTIONS + FORMATS + EVENT_TYPES))
    if request.method == "POST":
        action = request.form.get("action", "create")
        if action == "toggle":
            template_id = request.form.get("template_id", type=int)
            row = db.execute("SELECT * FROM task_templates WHERE id = ?", (template_id,)).fetchone()
            if row:
                db.execute("UPDATE task_templates SET active = ?, updated_at = ? WHERE id = ?", (0 if row["active"] else 1, now_iso(), template_id))
                db.commit()
                flash("Modelo ativado/inativado.", "success")
            return redirect(url_for("checklist_templates"))
        title = request.form.get("title", "").strip()
        trigger_type = request.form.get("trigger_type", "always")
        trigger_value = request.form.get("trigger_value", "").strip() or None
        try:
            due_offset = int(request.form.get("due_offset", "-7"))
        except ValueError:
            due_offset = -7
        if not title:
            flash("Informe o nome da tarefa do modelo.", "error")
        elif trigger_type not in trigger_types:
            flash("Condição de ativação inválida.", "error")
        elif trigger_type != "always" and not trigger_value:
            flash("Informe o valor que ativará esta tarefa.", "error")
        elif trigger_type == "support" and trigger_value not in SUPPORT_OPTIONS:
            flash("Selecione um serviço de apoio válido.", "error")
        elif trigger_type == "item" and trigger_value not in REQUESTED_ITEM_OPTIONS:
            flash("Selecione um item válido do formulário.", "error")
        elif trigger_type == "mode" and trigger_value not in FORMATS:
            flash("Selecione uma modalidade válida.", "error")
        elif trigger_type == "category" and trigger_value not in EVENT_TYPES:
            flash("Selecione uma categoria de evento válida.", "error")
        else:
            values = (title[:240], request.form.get("category", "Planejamento")[:80],
                      request.form.get("description", "").strip()[:1500] or None,
                      trigger_type, trigger_value, max(-365, min(365, due_offset)),
                      request.form.get("assigned_to", "").strip()[:180] or None)
            if action == "update":
                template_id = request.form.get("template_id", type=int)
                db.execute("""UPDATE task_templates SET title=?, category=?, description=?, trigger_type=?, trigger_value=?,
                           due_offset=?, assigned_to=?, updated_at=? WHERE id=?""", (*values, now_iso(), template_id))
                flash("Modelo de tarefa atualizado. A alteração vale para checklists gerados daqui em diante.", "success")
            else:
                order = db.execute("SELECT COALESCE(MAX(sort_order),0)+1 n FROM task_templates").fetchone()["n"]
                db.execute("""INSERT INTO task_templates(title, category, description, trigger_type, trigger_value,
                           due_offset, assigned_to, active, sort_order, created_at, updated_at)
                           VALUES (?, ?, ?, ?, ?, ?, ?, 1, ?, ?, ?)""", (*values, order, now_iso(), now_iso()))
                flash("Modelo de tarefa criado. Ele será aplicado a novas solicitações aprovadas.", "success")
            db.commit()
            return redirect(url_for("checklist_templates"))
    templates = db.execute("SELECT * FROM task_templates ORDER BY active DESC, sort_order, id").fetchall()
    return render_template("checklist_templates.html", templates=templates, categories=categories, trigger_types=trigger_types, trigger_values=trigger_values)


@app.route("/cadastros-apoio", methods=["GET", "POST"])
@login_required
def support_catalog_management():
    if session.get("user_role") != "admin":
        abort(404)
    db = get_db()
    categories = ["Materiais", "Infraestrutura", "Comunicação", "Credenciamento", "Logística", "Fornecedores", "Outro"]
    if request.method == "POST":
        action = request.form.get("action", "create")
        if action == "toggle":
            catalog_id = request.form.get("catalog_id", type=int)
            row = db.execute("SELECT * FROM support_catalog WHERE id = ?", (catalog_id,)).fetchone()
            if row:
                db.execute("UPDATE support_catalog SET active = ?, updated_at = ? WHERE id = ?", (0 if row["active"] else 1, now_iso(), catalog_id))
                db.commit()
                flash("Cadastro ativado/inativado.", "success")
            return redirect(url_for("support_catalog_management"))
        name = request.form.get("name", "").strip()
        status = request.form.get("status", "Precisa confirmar")
        try:
            quantity = int(request.form.get("quantity")) if (request.form.get("quantity") or "").strip() else None
            if quantity is not None and quantity < 0:
                raise ValueError
        except ValueError:
            quantity = -1
        if not name:
            flash("Informe o nome do item ou recurso.", "error")
        elif status not in SUPPORT_CATALOG_STATUSES:
            flash("Status do cadastro inválido.", "error")
        elif quantity == -1:
            flash("A quantidade deve ser um número igual ou maior que zero.", "error")
        else:
            values = (name[:180], request.form.get("category", "Outro")[:80], quantity,
                      request.form.get("location", "").strip()[:240] or None, status,
                      request.form.get("notes", "").strip()[:2000] or None)
            if action == "update":
                catalog_id = request.form.get("catalog_id", type=int)
                db.execute("""UPDATE support_catalog SET name=?, category=?, quantity=?, location=?, status=?, notes=?, updated_at=? WHERE id=?""",
                           (*values, now_iso(), catalog_id))
                flash("Cadastro atualizado.", "success")
            else:
                db.execute("""INSERT INTO support_catalog(name, category, quantity, location, status, notes, active, updated_at)
                           VALUES (?, ?, ?, ?, ?, ?, 1, ?)""", (*values, now_iso()))
                flash("Item adicionado ao cadastro de apoio.", "success")
            db.commit()
            return redirect(url_for("support_catalog_management"))
    records = db.execute("SELECT * FROM support_catalog ORDER BY active DESC, category, name").fetchall()
    return render_template("support_catalog.html", records=records, categories=categories)



@app.route("/equipe", methods=["GET", "POST"])
@login_required
def team_management():
    if session.get("user_role") != "admin":
        abort(404)
    db = get_db()
    if request.method == "POST":
        action = (request.form.get("action") or "create").strip()
        target_id = request.form.get("user_id", type=int)
        target = db.execute("SELECT * FROM users WHERE id = ?", (target_id,)).fetchone() if target_id else None

        if action == "create":
            email = request.form.get("email", "").strip().lower()
            password = request.form.get("password", "")
            role = request.form.get("role", "staff")
            if "@" not in email or "." not in email.split("@")[-1]:
                flash("Informe um e-mail válido.", "error")
            elif len(password) < 6:
                flash("A senha deve ter pelo menos 6 caracteres.", "error")
            elif role not in {"staff", "admin"}:
                flash("Perfil de acesso inválido.", "error")
            else:
                try:
                    db.execute("INSERT INTO users(email, password_hash, role, is_active, created_at) VALUES (?, ?, ?, 1, ?)",
                               (email, generate_password_hash(password), role, now_iso()))
                    db.commit()
                    flash("Acesso criado. No primeiro login, o perfil Equipe abre diretamente Triagem e acompanhamento.", "success")
                    return redirect(url_for("team_management"))
                except sqlite3.IntegrityError:
                    flash("Já existe um usuário cadastrado com esse e-mail.", "error")

        elif action == "update":
            if not target:
                flash("Usuário não encontrado.", "error")
            else:
                email = (request.form.get("email") or "").strip().lower()
                role = request.form.get("role", "staff")
                if "@" not in email or "." not in email.split("@")[-1]:
                    flash("Informe um e-mail válido.", "error")
                elif role not in {"staff", "admin"}:
                    flash("Perfil de acesso inválido.", "error")
                elif target["id"] == session.get("user_id") and role != "admin":
                    flash("Você não pode remover seu próprio perfil de administrador. Outro administrador deve realizar essa alteração.", "error")
                elif target["role"] == "admin" and target["is_active"] and role != "admin" and db.execute("SELECT COUNT(*) FROM users WHERE role='admin' AND is_active=1").fetchone()[0] <= 1:
                    flash("Não é possível retirar o perfil do último administrador ativo.", "error")
                else:
                    try:
                        db.execute("UPDATE users SET email=?, role=? WHERE id=?", (email, role, target_id))
                        db.commit()
                        if target_id == session.get("user_id"):
                            session["user_email"] = email
                            session["user_role"] = role
                        flash("Cadastro do usuário atualizado.", "success")
                    except sqlite3.IntegrityError:
                        flash("Já existe outro usuário com esse e-mail.", "error")

        elif action == "change_password":
            password = request.form.get("new_password", "")
            confirmation = request.form.get("confirm_password", "")
            if not target:
                flash("Usuário não encontrado.", "error")
            elif len(password) < 6:
                flash("A senha deve ter pelo menos 6 caracteres.", "error")
            elif password != confirmation:
                flash("A confirmação da senha não corresponde.", "error")
            else:
                db.execute("UPDATE users SET password_hash=? WHERE id=?", (generate_password_hash(password), target_id))
                db.commit()
                flash("Senha redefinida. Compartilhe a nova senha por um canal seguro.", "success")

        elif action == "toggle_active":
            if not target:
                flash("Usuário não encontrado.", "error")
            elif target_id == session.get("user_id"):
                flash("Você não pode bloquear sua própria conta enquanto está usando o sistema.", "error")
            elif target["role"] == "admin" and target["is_active"] and db.execute("SELECT COUNT(*) FROM users WHERE role='admin' AND is_active=1").fetchone()[0] <= 1:
                flash("Não é possível bloquear o último administrador ativo.", "error")
            else:
                new_value = 0 if target["is_active"] else 1
                db.execute("UPDATE users SET is_active=? WHERE id=?", (new_value, target_id))
                db.commit()
                flash("Conta bloqueada." if not new_value else "Conta desbloqueada.", "success")

        elif action == "delete":
            if not target:
                flash("Usuário não encontrado.", "error")
            elif target_id == session.get("user_id"):
                flash("Você não pode excluir sua própria conta enquanto está usando o sistema.", "error")
            elif target["role"] == "admin" and target["is_active"] and db.execute("SELECT COUNT(*) FROM users WHERE role='admin' AND is_active=1").fetchone()[0] <= 1:
                flash("Não é possível excluir o último administrador ativo.", "error")
            else:
                db.execute("DELETE FROM users WHERE id=?", (target_id,))
                db.commit()
                flash("Conta excluída.", "success")
        else:
            flash("Ação de administração não reconhecida.", "error")
        return redirect(url_for("team_management"))

    users = db.execute("SELECT id, email, role, is_active, created_at FROM users ORDER BY is_active DESC, created_at, email").fetchall()
    active_admins = db.execute("SELECT COUNT(*) FROM users WHERE role='admin' AND is_active=1").fetchone()[0]
    return render_template("team.html", users=users, active_admins=active_admins)



def parse_agenda_workbook(file_bytes: bytes, filename: str) -> dict:
    """Parse the FMABC monthly schedule layout and keep only event rows dated today or later."""
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.datetime import from_excel
    except ImportError as exc:
        raise RuntimeError("A biblioteca de leitura XLSX não está instalada. Instale as dependências do requirements.txt.") from exc

    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    today = date.today()
    imported = []
    skipped_past = 0

    def cell_text(value):
        if value is None:
            return ""
        return re.sub(r"\s+", " ", str(value)).strip()

    def parse_day(value):
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        if isinstance(value, (int, float)):
            dt = from_excel(value, workbook.epoch)
            if isinstance(dt, datetime):
                return dt.date()
            if isinstance(dt, date):
                return dt
        text_value = cell_text(value)
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                return datetime.strptime(text_value, fmt).date()
            except ValueError:
                pass
        return None

    def finalize(current):
        nonlocal skipped_past
        if not current:
            return
        event_date = current.pop("_parsed_date", None)
        if not event_date or event_date < today:
            skipped_past += 1
            return
        current["start_date"] = event_date.isoformat()
        current["title"] = current["title"].strip()
        if current["title"]:
            imported.append(current)

    for worksheet in workbook.worksheets:
        current = None
        for row_number, row in enumerate(worksheet.iter_rows(min_row=2, max_col=7, values_only=True), start=2):
            values = list(row) + [None] * max(0, 7 - len(row))
            event_date = parse_day(values[0])
            if event_date:
                finalize(current)
                current = {
                    "source_sheet": worksheet.title,
                    "source_row": row_number,
                    "_parsed_date": event_date,
                    "start_time": cell_text(values[1]),
                    "title": cell_text(values[2]),
                    "organizer": cell_text(values[3]),
                    "legacy_contact": cell_text(values[4]),
                    "venue": cell_text(values[5]),
                    "raw_notes": [],
                }
                if cell_text(values[6]):
                    current["raw_notes"].append(cell_text(values[6]))
            else:
                note = cell_text(values[6])
                if current and note:
                    current["raw_notes"].append(note)
        finalize(current)
    workbook.close()

    for event in imported:
        area_notes = []
        other_notes = []
        for raw in event["raw_notes"]:
            match = re.match(r"^\s*(Imprensa|Inspetoria|Limpeza|Manutenção|TI)\s*-\s*(.*)$", raw, re.I)
            if not match:
                other_notes.append(raw)
                continue
            area = match.group(1).strip().title()
            if area.lower() == "ti":
                area = "TI"
            details = match.group(2).strip()
            lowered = details.casefold()
            if "para conhecimento" in lowered:
                involvement, clean_details = "conhecimento", ""
            elif lowered == "psc" or lowered.startswith("psc") or lowered.endswith(" psc"):
                involvement, clean_details = "conhecimento", "Observação original: PSC (sigla preservada sem interpretação)."
            elif details:
                involvement, clean_details = "acao", details
            else:
                involvement, clean_details = "conhecimento", ""
            area_notes.append({"department": area, "involvement": involvement, "details": clean_details, "original": raw})
        event["department_notes"] = area_notes
        event["other_notes"] = other_notes
        event.pop("raw_notes", None)
        event["source_filename"] = filename[:240]
    return {"source_filename": filename[:240], "prepared_at": now_iso(), "today_filter": today.isoformat(), "skipped_past": skipped_past, "events": imported}


def import_agenda_payload(payload: dict) -> tuple[int, int, int, int]:
    db = get_db()
    inserted_events = skipped_events = inserted_department_links = task_count = 0
    departments = {row["name"].casefold(): row["id"] for row in db.execute("SELECT id, name FROM departments WHERE active = 1").fetchall()}
    for item in payload.get("events", []):
        existing = db.execute("SELECT id FROM events WHERE start_date=? AND lower(trim(title))=lower(trim(?)) LIMIT 1", (item["start_date"], item["title"])).fetchone()
        if existing:
            skipped_events += 1
            continue
        contact = (item.get("legacy_contact") or "").strip()
        digits = re.sub(r"\D", "", contact)
        contact_is_email = bool(re.match(r"^[^\s@]+@[^\s@]+\.[^\s@]+$", contact))
        contact_is_phone = (len(digits) >= 8 and not contact_is_email and not re.search(r"[A-Za-zÀ-ÿ]", contact))
        organizer_email = contact.lower() if contact_is_email else ""
        organizer_phone = contact if contact_is_phone else ""
        organizer = (item.get("organizer") or "").strip() or "Não informado — confirmar"
        venue = (item.get("venue") or "").strip() or None
        title_lower = (item.get("title") or "").casefold()
        if "colação" in title_lower or "formatura" in title_lower:
            event_type = "Cerimônia / Colação de grau"
        elif "simpósio" in title_lower or "congresso" in title_lower:
            event_type = "Congresso / Simpósio"
        elif "palestra" in title_lower or "mesa-redonda" in title_lower:
            event_type = "Palestra / Mesa-redonda"
        elif "curso" in title_lower or "oficina" in title_lower:
            event_type = "Curso / Oficina"
        else:
            event_type = "Evento institucional"
        format_source = f"{title_lower} {venue or ''}".casefold()
        event_format = "On-line" if "online" in format_source or "on-line" in format_source else ("Híbrido" if "híbrido" in format_source else "Presencial")
        active_supports = []
        for note in item.get("department_notes", []):
            if note.get("involvement") == "acao" and note.get("department") in SUPPORT_OPTIONS and note["department"] not in active_supports:
                active_supports.append(note["department"])
        other_notes = list(item.get("other_notes") or [])
        if contact and not (contact_is_email or contact_is_phone):
            other_notes.append(f"Contato conforme coluna original da planilha (tipo não identificado): {contact}")
        source_note = f"Importado da planilha {item.get('source_filename') or payload.get('source_filename') or 'agenda histórica'}, aba {item.get('source_sheet', '—')}, linha {item.get('source_row', '—')}. Dados ausentes precisam ser confirmados antes de novos comunicados."
        additional = "\n".join([source_note] + (["Observações originais da planilha: " + "; ".join(other_notes)] if other_notes else []))
        now = now_iso()
        cur = db.execute(
            """INSERT INTO events(protocol, title, department, organizer_name, organizer_email, organizer_phone,
               event_type, event_format, start_date, end_date, start_time, end_time, expected_attendees, venue,
               audience, description, support_json, event_modes_json, venues_json, requested_items_json,
               contact_email, legacy_contact_reference, authorized_by, cost_center, additional_info, internal_notes, event_type_other,
               external_guests, accessibility, status, priority, created_at, updated_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, NULL, 0, ?, ?, ?, ?, ?, ?, '[]', '', ?, NULL, NULL, ?, ?, NULL,
                       'Não informado', 'Não informado', 'em_planejamento', 'Normal', ?, ?)""",
            (f"FMABC-IMP-{item['start_date'].replace('-', '')}-{secrets.token_hex(2).upper()}", item["title"],
             "Agenda histórica importada", organizer, organizer_email, organizer_phone, event_type, event_format,
             item["start_date"], item["start_date"], item.get("start_time") or None, venue,
             "A confirmar", "Registro importado da agenda histórica. Confirmar dados incompletos com o Setor de Eventos.",
             json.dumps(active_supports, ensure_ascii=False), json.dumps([event_format], ensure_ascii=False),
             json.dumps([venue] if venue else [], ensure_ascii=False), contact or None, additional,
             f"Fonte: {item.get('source_sheet','—')} linha {item.get('source_row','—')}. Estado de importação: revisar.", now, now),
        )
        event_id = cur.lastrowid
        inserted_events += 1
        for note in item.get("department_notes", []):
            dept_id = departments.get((note.get("department") or "").casefold())
            if not dept_id:
                continue
            db.execute(
                "INSERT OR IGNORE INTO event_department_updates(event_id, department_id, involvement, details, notes, updated_by, updated_at) VALUES (?, ?, ?, ?, ?, ?, ?)",
                (event_id, dept_id, note.get("involvement", "conhecimento"), note.get("details") or None,
                 f"Texto original da planilha: {note.get('original','')}", session.get("user_email"), now),
            )
            inserted_department_links += 1
        event = db.execute("SELECT * FROM events WHERE id=?", (event_id,)).fetchone()
        task_count += add_checklist(event)
    db.commit()
    return inserted_events, skipped_events, inserted_department_links, task_count




@app.get("/comunicacao")
@login_required
def communication_center():
    db = get_db()
    today = date.today()
    try:
        year = int(request.args.get("year", str(today.year)))
        month = int(request.args.get("month", str(today.month)))
        if year < 2000 or year > 2100 or month < 1 or month > 12:
            raise ValueError
    except ValueError:
        year, month = today.year, today.month

    departments = db.execute("SELECT * FROM departments WHERE active = 1 ORDER BY name").fetchall()
    # Área selecionada para editar as providências por evento. É diferente da
    # lista múltipla de destinatários usados para gerar os comunicados.
    selected_id = request.args.get("edit_department_id", type=int)
    if selected_id is None:
        selected_id = request.args.get("department_id", type=int)  # compatibilidade com links antigos
    if selected_id and not any(d["id"] == selected_id for d in departments):
        selected_id = None
    if selected_id is None and departments:
        selected_id = departments[0]["id"]
    department = next((d for d in departments if d["id"] == selected_id), None)

    selected_recipient_ids = set(request.args.getlist("recipient_department_id", type=int))
    active_ids = {d["id"] for d in departments}
    selected_recipient_ids.intersection_update(active_ids)

    month_prefix = f"{year:04d}-{month:02d}"
    event_rows = db.execute(
        "SELECT * FROM events WHERE substr(start_date, 1, 7) = ? AND status IN ('aprovado', 'em_planejamento') ORDER BY start_date, start_time, title",
        (month_prefix,),
    ).fetchall()
    updates = {}
    if department:
        updates = {r["event_id"]: r for r in db.execute(
            "SELECT * FROM event_department_updates WHERE department_id = ?", (department["id"],)
        ).fetchall()}
    events = []
    for event in event_rows:
        item = dict(event)
        update = updates.get(event["id"])
        item["involvement"] = update["involvement"] if update else "conhecimento"
        item["details"] = update["details"] if update else ""
        item["update_notes"] = update["notes"] if update else ""
        item["update_id"] = update["id"] if update else None
        item["response_status"] = update["response_status"] if update and "response_status" in update.keys() else "aguardando"
        item["response_notes"] = update["response_notes"] if update and "response_notes" in update.keys() else None
        item["response_by"] = update["response_by"] if update and "response_by" in update.keys() else None
        item["responded_at"] = update["responded_at"] if update and "responded_at" in update.keys() else None
        item["receipt_confirmed_at"] = update["receipt_confirmed_at"] if update and "receipt_confirmed_at" in update.keys() else None
        item["receipt_confirmed_by"] = update["receipt_confirmed_by"] if update and "receipt_confirmed_by" in update.keys() else None
        item["area_responsible"] = update["area_responsible"] if update and "area_responsible" in update.keys() else None
        item["due_date"] = update["due_date"] if update and "due_date" in update.keys() else None
        item["task_status"] = update["task_status"] if update and "task_status" in update.keys() else "nao_iniciada"
        item["pending_reason"] = update["pending_reason"] if update and "pending_reason" in update.keys() else None
        item["response_token"] = update["response_token"] if update and "response_token" in update.keys() else None
        item["update_history"] = []
        if update:
            item["update_history"] = [dict(h) for h in db.execute("""SELECT changed_at,changed_by,source,action,change_summary
                FROM department_task_history WHERE update_id=? ORDER BY changed_at DESC,id DESC LIMIT 8""",(update["id"],)).fetchall()]
        events.append(item)

    # Pré-visualizações em rascunho do período ficam embutidas na mesma página
    # para facilitar a revisão e o envio em lote, sem abrir outra tela.
    draft_rows = db.execute(
        "SELECT b.*, d.name AS department_name, (SELECT COUNT(*) FROM communication_batch_events be WHERE be.batch_id=b.id) AS event_count "
        "FROM communication_batches b JOIN departments d ON d.id=b.department_id "
        "WHERE b.period_year=? AND b.period_month=? AND b.status='rascunho' ORDER BY b.created_at DESC, d.name",
        (year, month),
    ).fetchall()
    monthly_sent_count = db.execute(
        "SELECT COUNT(*) FROM communication_batches WHERE period_year=? AND period_month=? AND status='enviado'",
        (year, month),
    ).fetchone()[0]
    draft_batches = []
    for row in draft_rows:
        item = dict(row)
        item["items"] = db.execute(
            "SELECT * FROM communication_batch_events WHERE batch_id=? ORDER BY event_date, event_time, event_title",
            (item["id"],),
        ).fetchall()
        draft_batches.append(item)

    batches = db.execute(
        "SELECT b.*, d.name AS department_name FROM communication_batches b JOIN departments d ON d.id = b.department_id ORDER BY b.created_at DESC LIMIT 12"
    ).fetchall()
    years = list(range(max(2024, today.year - 2), min(2100, today.year + 6) + 1))
    all_departments = db.execute("SELECT * FROM departments ORDER BY active DESC, name").fetchall() if session.get("user_role") == "admin" else []
    return render_template(
        "communication.html", departments=departments, all_departments=all_departments,
        department=department, selected_id=selected_id, selected_recipient_ids=selected_recipient_ids,
        year=year, month=month, years=years, events=events, draft_batches=draft_batches,
        batches=batches, monthly_sent_count=monthly_sent_count,
        smtp_configured=smtp_is_configured(), roles=COMMUNICATION_ROLES,
        show_drafts=request.args.get("show_drafts") == "1",
    )


@app.post("/comunicacao/area")
@login_required
def save_communication_department():
    if session.get("user_role") != "admin":
        abort(404)
    db = get_db()
    action = request.form.get("action", "save")
    department_id = request.form.get("department_id", type=int)
    if action == "toggle":
        row = db.execute("SELECT * FROM departments WHERE id = ?", (department_id,)).fetchone()
        if row:
            db.execute("UPDATE departments SET active = ?, updated_at = ? WHERE id = ?", (0 if row["active"] else 1, now_iso(), department_id))
            db.commit()
            flash("Área ativada/inativada.", "success")
        return redirect(url_for("communication_center"))
    name = request.form.get("name", "").strip()
    email = request.form.get("recipient_email", "").strip().lower() or None
    if not name:
        flash("Informe o nome do departamento.", "error")
    elif len(name) > 100:
        flash("O nome do departamento é muito longo.", "error")
    elif email and ("@" not in email or "." not in email.split("@")[-1]):
        flash("Informe um e-mail válido para a área.", "error")
    else:
        try:
            if action == "update" and department_id:
                db.execute("UPDATE departments SET name = ?, recipient_email = ?, updated_at = ? WHERE id = ?", (name, email, now_iso(), department_id))
                flash("Área e destinatário atualizados.", "success")
            else:
                db.execute("INSERT INTO departments(name, recipient_email, active, created_at, updated_at) VALUES (?, ?, 1, ?, ?)", (name, email, now_iso(), now_iso()))
                flash("Área cadastrada.", "success")
            db.commit()
        except sqlite3.IntegrityError:
            flash("Já existe uma área com esse nome.", "error")
    return redirect(url_for("communication_center"))


@app.post("/comunicacao/vincular")
@login_required
def save_event_department_update():
    db = get_db()
    event_id = request.form.get("event_id", type=int)
    department_id = request.form.get("department_id", type=int)
    involvement = (request.form.get("involvement") or "conhecimento").strip()
    details = (request.form.get("details") or "").strip()[:2000] or None
    notes = (request.form.get("notes") or "").strip()[:2000] or None
    area_responsible = (request.form.get("area_responsible") or "").strip()[:120] or None
    due_date = (request.form.get("due_date") or "").strip() or None
    if due_date:
        try:
            date.fromisoformat(due_date)
        except ValueError:
            flash("Informe uma data válida para o prazo da providência.", "error")
            return redirect(url_for("communication_center"))

    event = get_event_or_404(event_id)
    dept = db.execute("SELECT * FROM departments WHERE id = ? AND active = 1", (department_id,)).fetchone()
    if not dept or involvement not in COMMUNICATION_ROLES:
        flash("Área ou tipo de participação inválido.", "error")
    else:
        old_row = db.execute("SELECT * FROM event_department_updates WHERE event_id=? AND department_id=?", (event_id,department_id)).fetchone()
        old = _department_task_snapshot(old_row)
        now = now_iso()
        reset_response = bool(old_row and (old_row["involvement"] != involvement or (old_row["details"] or "") != (details or "")))
        if old_row:
            response_status = "aguardando" if reset_response else old_row["response_status"]
            response_notes = None if reset_response else old_row["response_notes"]
            response_by = None if reset_response else old_row["response_by"]
            responded_at = None if reset_response else old_row["responded_at"]
            receipt_at = None if reset_response else old_row["receipt_confirmed_at"]
            receipt_by = None if reset_response else old_row["receipt_confirmed_by"]
            task_status = "nao_iniciada" if reset_response else old_row["task_status"]
            task_updated_at = None if reset_response else old_row["task_updated_at"]
            task_updated_by = None if reset_response else old_row["task_updated_by"]
            pending_reason = None if reset_response else old_row["pending_reason"]
            response_token = (old_row["response_token"] or secrets.token_urlsafe(32)) if involvement == "acao" else None
            db.execute("""UPDATE event_department_updates SET involvement=?,details=?,notes=?,updated_by=?,updated_at=?,
                response_status=?,response_notes=?,response_by=?,responded_at=?,receipt_confirmed_at=?,receipt_confirmed_by=?,
                task_status=?,task_updated_at=?,task_updated_by=?,pending_reason=?,area_responsible=?,due_date=?,response_token=?
                WHERE id=?""",
                (involvement,details,notes,session.get("user_email"),now,response_status,response_notes,response_by,responded_at,
                 receipt_at,receipt_by,task_status,task_updated_at,task_updated_by,pending_reason,area_responsible,due_date,response_token,old_row["id"]))
            update_id = old_row["id"]
            action = "Participação/providências atualizadas"
        else:
            response_token = secrets.token_urlsafe(32) if involvement == "acao" else None
            cursor = db.execute("""INSERT INTO event_department_updates
                (event_id,department_id,involvement,details,notes,updated_by,updated_at,response_status,response_token,
                 area_responsible,due_date,task_status)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,'nao_iniciada')""",
                (event_id,department_id,involvement,details,notes,session.get("user_email"),now,"aguardando",response_token,area_responsible,due_date))
            update_id = cursor.lastrowid
            action = "Providência criada"
        updated = db.execute("SELECT * FROM event_department_updates WHERE id=?",(update_id,)).fetchone()
        _record_department_task_history(db,update_id,event_id,department_id,old,_department_task_snapshot(updated),
            session.get("user_email"),"Setor de Eventos",action)
        db.commit()
        flash(f"Providência da área {dept['name']} atualizada. O acompanhamento e o histórico foram registrados.", "success")
    try:
        year, month = map(int, event["start_date"][:7].split("-"))
    except (ValueError, TypeError):
        year, month = date.today().year, date.today().month
    return redirect(url_for("communication_center", year=year, month=month, edit_department_id=department_id) + f"#event-{event_id}")


@app.post("/comunicacao/gerar")
@login_required
def generate_communication_batch():
    db = get_db()
    try:
        year = int(request.form.get("year", str(date.today().year)))
        month = int(request.form.get("month", str(date.today().month)))
        if year < 2000 or year > 2100 or not 1 <= month <= 12:
            raise ValueError
    except ValueError:
        flash("Selecione um mês e ano válidos.", "error")
        return redirect(url_for("communication_center"))

    # Aceita vários destinatários. O campo singular é mantido para links antigos.
    recipient_ids = request.form.getlist("recipient_department_ids", type=int)
    if not recipient_ids:
        legacy_id = request.form.get("department_id", type=int)
        if legacy_id:
            recipient_ids = [legacy_id]
    recipient_ids = list(dict.fromkeys(recipient_ids))
    if not recipient_ids:
        flash("Selecione pelo menos um departamento para receber a programação.", "error")
        return redirect(url_for("communication_center", year=year, month=month))

    created_count = 0
    created_events = 0
    skipped_departments = []
    for department_id in recipient_ids:
        dept = db.execute("SELECT * FROM departments WHERE id = ? AND active = 1", (department_id,)).fetchone()
        if not dept:
            continue
        period = f"{year:04d}-{month:02d}"
        event_rows = db.execute(
            "SELECT e.*, u.involvement AS area_involvement, u.details AS area_details, "
            "u.area_responsible AS area_responsible, u.due_date AS due_date "
            "FROM events e LEFT JOIN event_department_updates u ON u.event_id=e.id AND u.department_id=? "
            "WHERE substr(e.start_date,1,7)=? AND e.status IN ('aprovado','em_planejamento') "
            "ORDER BY e.start_date, e.start_time, e.title",
            (department_id, period),
        ).fetchall()
        relevant = [dict(r) for r in event_rows if (r["area_involvement"] or "conhecimento") != "excluir"]
        if not relevant:
            skipped_departments.append(dept["name"])
            continue
        for item in relevant:
            item["involvement"] = item.get("area_involvement") or "conhecimento"
            item["details"] = item.get("area_details") or ""
            item["response_url"] = ""
            if item["involvement"] == "acao":
                area_row = db.execute("SELECT id,response_token FROM event_department_updates WHERE event_id=? AND department_id=?", (item["id"], department_id)).fetchone()
                if area_row:
                    response_token = area_row["response_token"] or secrets.token_urlsafe(32)
                    db.execute("UPDATE event_department_updates SET response_token=?, response_status=CASE WHEN response_status IS NULL OR response_status='' THEN 'aguardando' ELSE response_status END WHERE id=?", (response_token, area_row["id"]))
                    base_url = _public_app_base_url()
                    if base_url:
                        item["response_url"] = base_url + url_for("department_response", token=response_token)
        subject, body = build_monthly_message(dept, year, month, relevant)
        cursor = db.execute(
            "INSERT INTO communication_batches(department_id, period_year, period_month, subject, body, recipient_email, status, created_by, created_at) "
            "VALUES (?, ?, ?, ?, ?, ?, 'rascunho', ?, ?)",
            (department_id, year, month, subject, body, dept["recipient_email"], session.get("user_email"), now_iso()),
        )
        batch_id = cursor.lastrowid
        for item in relevant:
            db.execute(
                "INSERT INTO communication_batch_events(batch_id, event_id, event_title, event_date, event_time, event_venue, involvement, details) "
                "VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                (batch_id, item["id"], item["title"], item["start_date"],
                 (f"{item['start_time']} às {item['end_time']}" if item.get("start_time") and item.get("end_time") else item.get("start_time")),
                 item["venue"], item["involvement"], item["details"]),
            )
        created_count += 1
        created_events += len(relevant)

    db.commit()
    if created_count:
        message = f"Foram criadas {created_count} prévia(s), totalizando {created_events} evento(s). Revise e envie na mesma página."
        if skipped_departments:
            message += " Sem eventos incluídos: " + ", ".join(skipped_departments) + "."
        flash(message, "success")
    else:
        flash("Não foi possível gerar prévias: não há eventos incluíveis nos departamentos selecionados para esse período.", "error")

    edit_department_id = request.form.get("edit_department_id", type=int)
    query = {"year": year, "month": month, "show_drafts": 1}
    if edit_department_id:
        query["edit_department_id"] = edit_department_id
    for dept_id in recipient_ids:
        query.setdefault("recipient_department_id", [])
        query["recipient_department_id"].append(dept_id)
    return redirect(url_for("communication_center", **query) + "#rascunhos-envio")


@app.get("/comunicacao/comunicado/<int:batch_id>")
@login_required
def communication_batch_detail(batch_id):
    db = get_db()
    batch = db.execute("SELECT b.*, d.name AS department_name FROM communication_batches b JOIN departments d ON d.id=b.department_id WHERE b.id=?", (batch_id,)).fetchone()
    if not batch:
        abort(404)
    items = db.execute("SELECT * FROM communication_batch_events WHERE batch_id = ? ORDER BY event_date, event_time, event_title", (batch_id,)).fetchall()
    return render_template("communication_batch.html", batch=batch, items=items, smtp_configured=smtp_is_configured())


def _communication_batch_redirect(batch):
    """Return to the monthly communication workspace when an inline action asks for it."""
    if request.form.get("return_to_center") == "1":
        return redirect(url_for(
            "communication_center", year=batch["period_year"], month=batch["period_month"],
            edit_department_id=batch["department_id"], show_drafts=1,
        ) + "#rascunhos-envio")
    return redirect(url_for("communication_batch_detail", batch_id=batch["id"]))


def _save_inline_batch_fields(db, batch):
    """Save editable draft fields submitted by the one-page communication workspace."""
    if "subject" not in request.form and "body" not in request.form and "recipient_email" not in request.form:
        return batch, True
    if batch["status"] != "rascunho":
        flash("Somente comunicados em rascunho podem ser alterados.", "error")
        return batch, False
    subject = request.form.get("subject", "").strip()[:250]
    body = request.form.get("body", "").strip()[:30000]
    recipient = request.form.get("recipient_email", "").strip().lower() or None
    if not subject or not body:
        flash("Informe o assunto e o conteúdo do comunicado.", "error")
        return batch, False
    if "\r" in subject or "\n" in subject or (recipient and ("\r" in recipient or "\n" in recipient)):
        flash("Não use quebras de linha no assunto ou no destinatário.", "error")
        return batch, False
    if recipient and ("@" not in recipient or "." not in recipient.split("@")[-1]):
        flash("O e-mail de destinatário não é válido.", "error")
        return batch, False
    db.execute("UPDATE communication_batches SET subject=?, body=?, recipient_email=? WHERE id=?", (subject, body, recipient, batch["id"]))
    db.commit()
    return db.execute("SELECT * FROM communication_batches WHERE id=?", (batch["id"],)).fetchone(), True


@app.post("/comunicacao/comunicado/<int:batch_id>/atualizar")
@login_required
def update_communication_batch(batch_id):
    db = get_db()
    batch = db.execute("SELECT * FROM communication_batches WHERE id = ?", (batch_id,)).fetchone()
    if not batch:
        abort(404)
    if batch["status"] != "rascunho":
        flash("Somente comunicados em rascunho podem ser editados.", "error")
    else:
        subject = request.form.get("subject", "").strip()[:250]
        body = request.form.get("body", "").strip()[:30000]
        recipient = request.form.get("recipient_email", "").strip().lower() or None
        if not subject or not body:
            flash("Informe o assunto e o conteúdo do comunicado.", "error")
        elif "\r" in subject or "\n" in subject or (recipient and ("\r" in recipient or "\n" in recipient)):
            flash("Não use quebras de linha no assunto ou no destinatário.", "error")
        elif recipient and ("@" not in recipient or "." not in recipient.split("@")[-1]):
            flash("O e-mail de destinatário não é válido.", "error")
        else:
            db.execute("UPDATE communication_batches SET subject=?, body=?, recipient_email=? WHERE id=?", (subject, body, recipient, batch_id))
            db.commit()
            flash("Prévia e destinatário salvos.", "success")
    return _communication_batch_redirect(batch)


def _mark_batch_sent(db, batch_id: int, method: str) -> None:
    db.execute("UPDATE communication_batches SET status='enviado', sent_at=?, sent_by=?, send_method=?, error_message=NULL WHERE id=?",
               (now_iso(), session.get("user_email"), method, batch_id))
    db.commit()


@app.post("/comunicacao/comunicado/<int:batch_id>/enviar")
@login_required
def send_communication_batch(batch_id):
    db = get_db()
    batch = db.execute("SELECT * FROM communication_batches WHERE id=?", (batch_id,)).fetchone()
    if not batch:
        abort(404)
    if batch["status"] != "rascunho":
        flash("Este comunicado já foi marcado como enviado e não pode ser reenviado por esta ação.", "error")
        return _communication_batch_redirect(batch)
    batch, fields_ok = _save_inline_batch_fields(db, batch)
    if not fields_ok:
        return _communication_batch_redirect(batch)
    recipient = (batch["recipient_email"] or "").strip()
    cfg = smtp_settings()
    if not recipient or not smtp_is_configured():
        flash("O envio automático ainda não está configurado. Cadastre o e-mail da área e configure SMTP no arquivo .env; ou envie pelo e-mail institucional e depois registre o envio manual.", "error")
        return _communication_batch_redirect(batch)
    if any(ch in recipient or ch in (batch["subject"] or "") for ch in ["\r", "\n"]):
        flash("O destinatário ou assunto contém caracteres não permitidos.", "error")
        return _communication_batch_redirect(batch)
    msg = EmailMessage()
    msg["Subject"] = batch["subject"]
    msg["From"] = cfg["from_email"]
    msg["To"] = recipient
    msg.set_content(batch["body"])
    try:
        with smtplib.SMTP(cfg["host"], cfg["port"], timeout=20) as server:
            if cfg["use_tls"]:
                server.starttls()
            if cfg["user"]:
                server.login(cfg["user"], cfg["password"])
            server.send_message(msg)
        _mark_batch_sent(db, batch_id, "smtp")
        flash(f"Comunicado enviado para {recipient} e registrado no histórico.", "success")
    except Exception as exc:
        safe_error = f"Falha no envio SMTP: {type(exc).__name__}"
        db.execute("UPDATE communication_batches SET error_message=? WHERE id=?", (safe_error, batch_id))
        db.commit()
        flash("Não foi possível enviar o comunicado. Confira a configuração com a TI; o rascunho foi preservado.", "error")
    return _communication_batch_redirect(batch)


@app.post("/comunicacao/comunicado/<int:batch_id>/registrar-envio-manual")
@login_required
def register_manual_communication_batch(batch_id):
    db = get_db()
    batch = db.execute("SELECT * FROM communication_batches WHERE id=?", (batch_id,)).fetchone()
    if not batch:
        abort(404)
    if batch["status"] != "rascunho":
        flash("Este comunicado já está registrado como enviado.", "error")
    else:
        batch, fields_ok = _save_inline_batch_fields(db, batch)
        if not fields_ok:
            return _communication_batch_redirect(batch)
        if not (batch["recipient_email"] or "").strip():
            flash("Antes de registrar o envio manual, informe o e-mail destinatário.", "error")
        elif request.form.get("confirm_sent") != "yes":
            flash("Marque a confirmação de que o e-mail foi enviado fora do sistema.", "error")
        else:
            _mark_batch_sent(db, batch_id, "manual")
            flash("Envio manual registrado no histórico.", "success")
    return _communication_batch_redirect(batch)


@app.post("/comunicacao/importar-agenda")
@login_required
def preview_agenda_import():
    if session.get("user_role") != "admin":
        abort(404)
    upload = request.files.get("agenda_file")
    if not upload or not upload.filename:
        flash("Selecione a planilha .xlsx para importar.", "error")
        return redirect(url_for("communication_center"))
    filename = Path(upload.filename).name
    if not filename.lower().endswith(".xlsx"):
        flash("Envie um arquivo .xlsx do Excel.", "error")
        return redirect(url_for("communication_center"))
    try:
        payload = parse_agenda_workbook(upload.read(), filename)
    except Exception as exc:
        flash(f"Não foi possível ler a planilha. Confira se o arquivo é um .xlsx válido ({type(exc).__name__}).", "error")
        return redirect(url_for("communication_center"))
    if not payload["events"]:
        flash(f"A planilha não contém eventos com data a partir de {payload['today_filter']} para importar.", "error")
        return redirect(url_for("communication_center"))
    token = secrets.token_urlsafe(24)
    preview_dir = INSTANCE_DIR / "import_previews"
    preview_dir.mkdir(exist_ok=True)
    (preview_dir / f"{token}.json").write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return render_template("agenda_import_preview.html", payload=payload, token=token)


@app.post("/comunicacao/importar-agenda/confirmar")
@login_required
def confirm_agenda_import():
    if session.get("user_role") != "admin":
        abort(404)
    token = request.form.get("token", "")
    if not re.fullmatch(r"[A-Za-z0-9_-]{20,80}", token):
        abort(400, description="Identificador de prévia inválido.")
    preview_path = INSTANCE_DIR / "import_previews" / f"{token}.json"
    if not preview_path.exists():
        flash("A prévia expirou ou já foi utilizada. Envie a planilha novamente.", "error")
        return redirect(url_for("communication_center"))
    try:
        payload = json.loads(preview_path.read_text(encoding="utf-8"))
        inserted, skipped, links, tasks = import_agenda_payload(payload)
        preview_path.unlink(missing_ok=True)
        flash(f"Importação concluída: {inserted} evento(s) inserido(s), {skipped} duplicado(s) ignorado(s), {links} vínculos com áreas e {tasks} tarefas de checklist criadas.", "success")
    except Exception as exc:
        get_db().rollback()
        flash(f"A importação não foi concluída; a prévia foi preservada. Erro: {type(exc).__name__}.", "error")
    return redirect(url_for("communication_center"))


@app.get("/health")
def health_check():
    """Health endpoint for hosting-platform checks."""
    try:
        db = get_db()
        db.execute("SELECT 1").fetchone()
        return jsonify({"status": "ok"}), 200
    except Exception:
        app.logger.exception("Health check failed")
        return jsonify({"status": "error"}), 503


@app.errorhandler(400)
def bad_request(error):
    return render_template("error.html", code=400, message=getattr(error, "description", "Não foi possível processar a solicitação.")), 400


@app.errorhandler(404)
def not_found(_error):
    return render_template("error.html", code=404, message="A página ou evento solicitado não foi encontrado."), 404


with app.app_context():
    init_db()

if __name__ == "__main__":
    app.run(host=os.getenv("HOST", "127.0.0.1"), port=int(os.getenv("PORT", "5000")), debug=os.getenv("DEBUG", "0") == "1")
