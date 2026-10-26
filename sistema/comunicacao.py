"""Agenda e tarefas por área, comunicados mensais e importação da planilha de agenda."""
from __future__ import annotations

from .core import *  # noqa: F401,F403


def build_monthly_message(department, year: int, month: int, events: list[dict]) -> tuple[str, str]:
    month_names = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro", "novembro", "dezembro"]
    subject = f"Agenda de eventos FMABC — {month_names[month - 1].capitalize()} de {year} — {department['name']}"
    lines = [
        "Prezados,", "", 
        f"Encaminhamos a programação de eventos da FMABC prevista para {month_names[month - 1]} de {year}, para organização e conhecimento da área {department['name']}.",
        "As indicações de envolvimento e as observações abaixo refletem os registros atuais do Setor de Eventos. Favor sinalizar eventuais ajustes.", ""
    ]
    for idx, item in enumerate(events, start=1):
        try:
            d = date.fromisoformat(item["start_date"]).strftime("%d/%m/%Y")
        except (TypeError, ValueError):
            d = item["start_date"] or "Data a confirmar"
        lines.append(f"{idx}. {item['title']}")
        time_text = (item.get("start_time") or "").strip()
        end_text = (item.get("end_time") or "").strip()
        if time_text and end_text and end_text not in time_text:
            time_text = f"{time_text} às {end_text}"
        lines.append(f"   Data: {d}" + (f" | Horário: {time_text}" if time_text else ""))
        lines.append(f"   Local: {item['venue'] or 'A confirmar'}")
        responsible = (item.get("organizer_name") or "").strip()
        if responsible:
            lines.append(f"   Responsável registrado: {responsible}")
        legacy_contact = (item.get("legacy_contact_reference") or "").strip()
        if legacy_contact:
            lines.append(f"   Contato/referência registrada na agenda: {legacy_contact}")
        else:
            contacts = []
            if (item.get("organizer_phone") or "").strip():
                contacts.append("Telefone/WhatsApp: " + item["organizer_phone"].strip())
            if (item.get("contact_email") or "").strip():
                contacts.append("E-mail de contato: " + item["contact_email"].strip())
            for contact_line in contacts:
                lines.append("   " + contact_line)
        lines.append(f"   Participação da área: {COMMUNICATION_ROLES.get(item['involvement'], item['involvement'])}")
        if item.get("area_responsible"):
            lines.append(f"   Responsável da área: {item['area_responsible']}")
        if item.get("due_date"):
            try:
                due_text = date.fromisoformat(item['due_date']).strftime("%d/%m/%Y")
            except (TypeError, ValueError):
                due_text = item['due_date']
            lines.append(f"   Prazo da providência: {due_text}")
        if item.get("details"):
            lines.append(f"   Providências solicitadas: {item['details']}")
        if item.get("response_url"):
            lines.append(f"   Confirme o recebimento, informe o responsável e atualize a situação da providência: {item['response_url']}")
        elif item.get("involvement") == "acao":
            lines.append("   Ação necessária: confirme o recebimento, indique o responsável da área e atualize a situação da providência.")
        lines.append("")
    lines += ["Pedimos que as áreas envolvidas organizem as providências indicadas e comuniquem o Setor de Eventos sobre possíveis impedimentos ou alterações.", "", "Atenciosamente,", "Setor de Eventos", "Centro Universitário FMABC"]
    return subject, "\n".join(lines)


def parse_agenda_workbook(file_bytes: bytes, filename: str) -> dict:
    """Parse the FMABC monthly schedule layout and keep only event rows dated today or later."""
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.datetime import from_excel
    except ImportError as exc:
        raise RuntimeError("A biblioteca de leitura XLSX não está instalada. Instale as dependências do requirements.txt.") from exc

    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    today = date.today()
    imported = []
    skipped_past = 0

    def cell_text(value):
        if value is None:
            return ""
        return re.sub(r"\s+", " ", str(value)).strip()

    def parse_day(value):
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        if isinstance(value, (int, float)):
            dt = from_excel(value, workbook.epoch)
            if isinstance(dt, datetime):
                return dt.date()
            if isinstance(dt, date):
                return dt
        text_value = cell_text(value)
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                return datetime.strptime(text_value, fmt).date()
            except ValueError:
                pass
        return None

    def finalize(current):
        nonlocal skipped_past
        if not current:
            return
        event_date = current.pop("_parsed_date", None)
        if not event_date or event_date < today:
            skipped_past += 1
            return
        current["start_date"] = event_date.isoformat()
        current["title"] = current["title"].strip()
        if current["title"]:
            imported.append(current)

    for worksheet in workbook.worksheets:
        current = None
        for row_number, row in enumerate(worksheet.iter_rows(min_row=2, max_col=7, values_only=True), start=2):
            values = list(row) + [None] * max(0, 7 - len(row))
            event_date = parse_day(values[0])
            if event_date:
                finalize(current)
                current = {
                    "source_sheet": worksheet.title,
                    "source_row": row_number,
                    "_parsed_date": event_date,
                    "start_time": cell_text(values[1]),
                    "title": cell_text(values[2]),
                    "organizer": cell_text(values[3]),
                    "legacy_contact": cell_text(values[4]),
                    "venue": cell_text(values[5]),
                    "raw_notes": [],
                }
                if cell_text(values[6]):
                    current["raw_notes"].append(cell_text(values[6]))
            else:
                note = cell_text(values[6])
                if current and note:
                    current["raw_notes"].append(note)
        finalize(current)
    workbook.close()

    for event in imported:
        area_notes = []
        other_notes = []
        for raw in event["raw_notes"]:
            match = re.match(r"^\s*(Imprensa|Inspetoria|Limpeza|Manutenção|TI)\s*-\s*(.*)$", raw, re.I)
            if not match:
                other_notes.append(raw)
                continue
            area = match.group(1).strip().title()
            if area.lower() == "ti":
                area = "TI"
            details = match.group(2).strip()
            lowered = details.casefold()
            if "para conhecimento" in lowered:
                involvement, clean_details = "conhecimento", ""
            elif lowered == "psc" or lowered.startswith("psc") or lowered.endswith(" psc"):
                involvement, clean_details = "conhecimento", "Observação original: PSC (sigla preservada sem interpretação)."
            elif details:
                involvement, clean_details = "acao", details
            else:
                involvement, clean_details = "conhecimento", ""
            area_notes.append({"department": area, "involvement": involvement, "details": clean_details, "original": raw})
        event["department_notes"] = area_notes
        event["other_notes"] = other_notes
        event.pop("raw_notes", None)
        event["source_filename"] = filename[:240]
    return {"source_filename": filename[:240], "prepared_at": now_iso(), "today_filter": today.isoformat(), "skipped_past": skipped_past, "events": imported}


def import_agenda_payload(payload: dict) -> tuple[int, int, int, int]:
    db = get_db()
    inserted_events = skipped_events = inserted_department_links = task_count = 0
    imported_ids = []
    departments = {row["name"].casefold(): row["id"] for row in db.execute("SELECT id, name FROM departments WHERE active = 1").fetchall()}
    for item in payload.get("events", []):
        existing = db.execute("SELECT id FROM events WHERE start_date=? AND lower(trim(title))=lower(trim(?)) LIMIT 1", (item["start_date"], item["title"])).fetchone()
        if existing:
            skipped_events += 1
            continue
        contact = (item.get("legacy_contact") or "").strip()
        digits = re.sub(r"\D", "", contact)
        contact_is_email = bool(re.match(r"^[^\s@]+@[^\s@]+\.[^\s@]+$", contact))
        contact_is_phone = (len(digits) >= 8 and not contact_is_email and not re.search(r"[A-Za-zÀ-ÿ]", contact))
        organizer_email = contact.lower() if contact_is_email else ""
        organizer_phone = contact if contact_is_phone else ""
        organizer = (item.get("organizer") or "").strip() or "Não informado — confirmar"
        venue = (item.get("venue") or "").strip() or None
        title_lower = (item.get("title") or "").casefold()
        if "colação" in title_lower or "formatura" in title_lower:
            event_type = "Cerimônia / Colação de grau"
        elif "simpósio" in title_lower or "congresso" in title_lower:
            event_type = "Congresso / Simpósio"
        elif "palestra" in title_lower or "mesa-redonda" in title_lower:
            event_type = "Palestra / Mesa-redonda"
        elif "curso" in title_lower or "oficina" in title_lower:
            event_type = "Curso / Oficina"
        else:
            event_type = "Evento institucional"
        format_source = f"{title_lower} {venue or ''}".casefold()
        event_format = "On-line" if "online" in format_source or "on-line" in format_source else ("Híbrido" if "híbrido" in format_source else "Presencial")
        active_supports = []
        for note in item.get("department_notes", []):
            if note.get("involvement") == "acao" and note.get("department") in SUPPORT_OPTIONS and note["department"] not in active_supports:
                active_supports.append(note["department"])
        other_notes = list(item.get("other_notes") or [])
        if contact and not (contact_is_email or contact_is_phone):
            other_notes.append(f"Contato conforme coluna original da planilha (tipo não identificado): {contact}")
        source_note = f"Importado da planilha {item.get('source_filename') or payload.get('source_filename') or 'agenda histórica'}, aba {item.get('source_sheet', '—')}, linha {item.get('source_row', '—')}. Dados ausentes precisam ser confirmados antes de novos comunicados."
        additional = "\n".join([source_note] + (["Observações originais da planilha: " + "; ".join(other_notes)] if other_notes else []))
        now = now_iso()
        cur = db.execute(
            """INSERT INTO events(protocol, title, department, organizer_name, organizer_email, organizer_phone,
               event_type, event_format, start_date, end_date, start_time, end_time, expected_attendees, venue,
               audience, description, support_json, event_modes_json, venues_json, requested_items_json,
               contact_email, legacy_contact_reference, authorized_by, cost_center, additional_info, internal_notes, event_type_other,
               external_guests, accessibility, status, priority, created_at, updated_at)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, NULL, 0, ?, ?, ?, ?, ?, ?, '[]', '', ?, NULL, NULL, ?, ?, NULL,
                       'Não informado', 'Não informado', 'em_planejamento', 'Normal', ?, ?)""",
            (f"FMABC-IMP-{item['start_date'].replace('-', '')}-{secrets.token_hex(2).upper()}", item["title"],
             "Agenda histórica importada", organizer, organizer_email, organizer_phone, event_type, event_format,
             item["start_date"], item["start_date"], item.get("start_time") or None, venue,
             "A confirmar", "Registro importado da agenda histórica. Confirmar dados incompletos com o Setor de Eventos.",
             json.dumps(active_supports, ensure_ascii=False), json.dumps([event_format], ensure_ascii=False),
             json.dumps([venue] if venue else [], ensure_ascii=False), contact or None, additional,
             f"Fonte: {item.get('source_sheet','—')} linha {item.get('source_row','—')}. Estado de importação: revisar.", now, now),
        )
        event_id = cur.lastrowid
        inserted_events += 1
        for note in item.get("department_notes", []):
            dept_id = departments.get((note.get("department") or "").casefold())
            if not dept_id:
                continue
            db.execute(
                "INSERT OR IGNORE INTO event_department_updates(event_id, department_id, involvement, details, notes, updated_by, updated_at) VALUES (?, ?, ?, ?, ?, ?, ?)",
                (event_id, dept_id, note.get("involvement", "conhecimento"), note.get("details") or None,
                 f"Texto original da planilha: {note.get('original','')}", session.get("user_email"), now),
            )
            inserted_department_links += 1
        event = db.execute("SELECT * FROM events WHERE id=?", (event_id,)).fetchone()
        task_count += add_checklist(event, commit=False)
        imported_ids.append(event_id)
    # Tudo ou nada: se qualquer linha falhar, a exceção desfaz a importação inteira.
    db.commit()
    payload["_imported_ids"] = imported_ids
    return inserted_events, skipped_events, inserted_department_links, task_count


@app.get("/comunicacao")
@login_required
def communication_center():
    db = get_db()
    today = date.today()
    try:
        year = int(request.args.get("year", str(today.year)))
        month = int(request.args.get("month", str(today.month)))
        if year < 2000 or year > 2100 or month < 1 or month > 12:
            raise ValueError
    except ValueError:
        year, month = today.year, today.month

    departments = db.execute("SELECT * FROM departments WHERE active = 1 ORDER BY name").fetchall()
    # Área selecionada para editar as providências por evento. É diferente da
    # lista múltipla de destinatários usados para gerar os comunicados.
    selected_id = request.args.get("edit_department_id", type=int)
    if selected_id is None:
        selected_id = request.args.get("department_id", type=int)  # compatibilidade com links antigos
    if selected_id and not any(d["id"] == selected_id for d in departments):
        selected_id = None
    if selected_id is None and departments:
        selected_id = departments[0]["id"]
    department = next((d for d in departments if d["id"] == selected_id), None)

    selected_recipient_ids = set(request.args.getlist("recipient_department_id", type=int))
    active_ids = {d["id"] for d in departments}
    selected_recipient_ids.intersection_update(active_ids)

    month_prefix = f"{year:04d}-{month:02d}"
    event_rows = db.execute(
        "SELECT * FROM events WHERE substr(start_date, 1, 7) = ? AND status IN ('aprovado', 'em_planejamento') ORDER BY start_date, start_time, title",
        (month_prefix,),
    ).fetchall()
    updates = {}
    if department:
        updates = {r["event_id"]: r for r in db.execute(
            "SELECT * FROM event_department_updates WHERE department_id = ?", (department["id"],)
        ).fetchall()}
    events = []
    for event in event_rows:
        item = dict(event)
        update = updates.get(event["id"])
        item["involvement"] = update["involvement"] if update else "conhecimento"
        item["details"] = update["details"] if update else ""
        item["update_notes"] = update["notes"] if update else ""
        item["update_id"] = update["id"] if update else None
        item["response_status"] = update["response_status"] if update and "response_status" in update.keys() else "aguardando"
        item["response_notes"] = update["response_notes"] if update and "response_notes" in update.keys() else None
        item["response_by"] = update["response_by"] if update and "response_by" in update.keys() else None
        item["responded_at"] = update["responded_at"] if update and "responded_at" in update.keys() else None
        item["receipt_confirmed_at"] = update["receipt_confirmed_at"] if update and "receipt_confirmed_at" in update.keys() else None
        item["receipt_confirmed_by"] = update["receipt_confirmed_by"] if update and "receipt_confirmed_by" in update.keys() else None
        item["area_responsible"] = update["area_responsible"] if update and "area_responsible" in update.keys() else None
        item["due_date"] = update["due_date"] if update and "due_date" in update.keys() else None
        item["task_status"] = update["task_status"] if update and "task_status" in update.keys() else "nao_iniciada"
        item["pending_reason"] = update["pending_reason"] if update and "pending_reason" in update.keys() else None
        item["response_token"] = update["response_token"] if update and "response_token" in update.keys() else None
        item["update_history"] = []
        if update:
            item["update_history"] = [dict(h) for h in db.execute("""SELECT changed_at,changed_by,source,action,change_summary
                FROM department_task_history WHERE update_id=? ORDER BY changed_at DESC,id DESC LIMIT 8""",(update["id"],)).fetchall()]
        events.append(item)

    # Pré-visualizações em rascunho do período ficam embutidas na mesma página
    # para facilitar a revisão e o envio em lote, sem abrir outra tela.
    draft_rows = db.execute(
        "SELECT b.*, d.name AS department_name, (SELECT COUNT(*) FROM communication_batch_events be WHERE be.batch_id=b.id) AS event_count "
        "FROM communication_batches b JOIN departments d ON d.id=b.department_id "
        "WHERE b.period_year=? AND b.period_month=? AND b.status='rascunho' ORDER BY b.created_at DESC, d.name",
        (year, month),
    ).fetchall()
    monthly_sent_count = db.execute(
        "SELECT COUNT(*) FROM communication_batches WHERE period_year=? AND period_month=? AND status='enviado'",
        (year, month),
    ).fetchone()[0]
    draft_batches = []
    for row in draft_rows:
        item = dict(row)
        item["items"] = db.execute(
            "SELECT * FROM communication_batch_events WHERE batch_id=? ORDER BY event_date, event_time, event_title",
            (item["id"],),
        ).fetchall()
        draft_batches.append(item)

    batches = db.execute(
        "SELECT b.*, d.name AS department_name FROM communication_batches b JOIN departments d ON d.id = b.department_id ORDER BY b.created_at DESC LIMIT 12"
    ).fetchall()
    years = list(range(max(2024, today.year - 2), min(2100, today.year + 6) + 1))
    all_departments = db.execute("SELECT * FROM departments ORDER BY active DESC, name").fetchall() if session.get("user_role") == "admin" else []
    return render_template(
        "communication.html", departments=departments, all_departments=all_departments,
        department=department, selected_id=selected_id, selected_recipient_ids=selected_recipient_ids,
        year=year, month=month, years=years, events=events, draft_batches=draft_batches,
        batches=batches, monthly_sent_count=monthly_sent_count,
        smtp_configured=smtp_is_configured(), roles=COMMUNICATION_ROLES,
        show_drafts=request.args.get("show_drafts") == "1",
    )


@app.post("/comunicacao/area")
@login_required
def save_communication_department():
    if session.get("user_role") != "admin":
        abort(404)
    db = get_db()
    action = request.form.get("action", "save")
    department_id = request.form.get("department_id", type=int)
    if action == "toggle":
        row = db.execute("SELECT * FROM departments WHERE id = ?", (department_id,)).fetchone()
        if row:
            db.execute("UPDATE departments SET active = ?, updated_at = ? WHERE id = ?", (0 if row["active"] else 1, now_iso(), department_id))
            db.commit()
            flash("Área ativada/inativada.", "success")
        return redirect(url_for("communication_center"))
    name = request.form.get("name", "").strip()
    email = request.form.get("recipient_email", "").strip().lower() or None
    if not name:
        flash("Informe o nome do departamento.", "error")
    elif len(name) > 100:
        flash("O nome do departamento é muito longo.", "error")
    elif email and ("@" not in email or "." not in email.split("@")[-1]):
        flash("Informe um e-mail válido para a área.", "error")
    else:
        try:
            if action == "update" and department_id:
                db.execute("UPDATE departments SET name = ?, recipient_email = ?, updated_at = ? WHERE id = ?", (name, email, now_iso(), department_id))
                flash("Área e destinatário atualizados.", "success")
            else:
                db.execute("INSERT INTO departments(name, recipient_email, active, created_at, updated_at) VALUES (?, ?, 1, ?, ?)", (name, email, now_iso(), now_iso()))
                flash("Área cadastrada.", "success")
            db.commit()
        except sqlite3.IntegrityError:
            flash("Já existe uma área com esse nome.", "error")
    return redirect(url_for("communication_center"))


@app.post("/comunicacao/vincular")
@login_required
def save_event_department_update():
    db = get_db()
    event_id = request.form.get("event_id", type=int)
    department_id = request.form.get("department_id", type=int)
    involvement = (request.form.get("involvement") or "conhecimento").strip()
    details = (request.form.get("details") or "").strip()[:2000] or None
    notes = (request.form.get("notes") or "").strip()[:2000] or None
    area_responsible = (request.form.get("area_responsible") or "").strip()[:120] or None
    due_date = (request.form.get("due_date") or "").strip() or None
    if due_date:
        try:
            date.fromisoformat(due_date)
        except ValueError:
            flash("Informe uma data válida para o prazo da providência.", "error")
            return redirect(url_for("communication_center"))

    event = get_event_or_404(event_id)
    dept = db.execute("SELECT * FROM departments WHERE id = ? AND active = 1", (department_id,)).fetchone()
    if not dept or involvement not in COMMUNICATION_ROLES:
        flash("Área ou tipo de participação inválido.", "error")
    else:
        old_row = db.execute("SELECT * FROM event_department_updates WHERE event_id=? AND department_id=?", (event_id,department_id)).fetchone()
        old = _department_task_snapshot(old_row)
        now = now_iso()
        reset_response = bool(old_row and (old_row["involvement"] != involvement or (old_row["details"] or "") != (details or "")))
        if old_row:
            response_status = "aguardando" if reset_response else old_row["response_status"]
            response_notes = None if reset_response else old_row["response_notes"]
            response_by = None if reset_response else old_row["response_by"]
            responded_at = None if reset_response else old_row["responded_at"]
            receipt_at = None if reset_response else old_row["receipt_confirmed_at"]
            receipt_by = None if reset_response else old_row["receipt_confirmed_by"]
            task_status = "nao_iniciada" if reset_response else old_row["task_status"]
            task_updated_at = None if reset_response else old_row["task_updated_at"]
            task_updated_by = None if reset_response else old_row["task_updated_by"]
            pending_reason = None if reset_response else old_row["pending_reason"]
            response_token = (old_row["response_token"] or secrets.token_urlsafe(32)) if involvement == "acao" else None
            db.execute("""UPDATE event_department_updates SET involvement=?,details=?,notes=?,updated_by=?,updated_at=?,
                response_status=?,response_notes=?,response_by=?,responded_at=?,receipt_confirmed_at=?,receipt_confirmed_by=?,
                task_status=?,task_updated_at=?,task_updated_by=?,pending_reason=?,area_responsible=?,due_date=?,response_token=?
                WHERE id=?""",
                (involvement,details,notes,session.get("user_email"),now,response_status,response_notes,response_by,responded_at,
                 receipt_at,receipt_by,task_status,task_updated_at,task_updated_by,pending_reason,area_responsible,due_date,response_token,old_row["id"]))
            update_id = old_row["id"]
            action = "Participação/providências atualizadas"
        else:
            response_token = secrets.token_urlsafe(32) if involvement == "acao" else None
            cursor = db.execute("""INSERT INTO event_department_updates
                (event_id,department_id,involvement,details,notes,updated_by,updated_at,response_status,response_token,
                 area_responsible,due_date,task_status)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,'nao_iniciada')""",
                (event_id,department_id,involvement,details,notes,session.get("user_email"),now,"aguardando",response_token,area_responsible,due_date))
            update_id = cursor.lastrowid
            action = "Providência criada"
        updated = db.execute("SELECT * FROM event_department_updates WHERE id=?",(update_id,)).fetchone()
        _record_department_task_history(db,update_id,event_id,department_id,old,_department_task_snapshot(updated),
            session.get("user_email"),"Setor de Eventos",action)
        db.commit()
        flash(f"Providência da área {dept['name']} atualizada. O acompanhamento e o histórico foram registrados.", "success")
    try:
        year, month = map(int, event["start_date"][:7].split("-"))
    except (ValueError, TypeError):
        year, month = date.today().year, date.today().month
    return redirect(url_for("communication_center", year=year, month=month, edit_department_id=department_id) + f"#event-{event_id}")


@app.post("/comunicacao/gerar")
@login_required
def generate_communication_batch():
    db = get_db()
    try:
        year = int(request.form.get("year", str(date.today().year)))
        month = int(request.form.get("month", str(date.today().month)))
        if year < 2000 or year > 2100 or not 1 <= month <= 12:
            raise ValueError
    except ValueError:
        flash("Selecione um mês e ano válidos.", "error")
        return redirect(url_for("communication_center"))

    # Aceita vários destinatários. O campo singular é mantido para links antigos.
    recipient_ids = request.form.getlist("recipient_department_ids", type=int)
    if not recipient_ids:
        legacy_id = request.form.get("department_id", type=int)
        if legacy_id:
            recipient_ids = [legacy_id]
    recipient_ids = list(dict.fromkeys(recipient_ids))
    if not recipient_ids:
        flash("Selecione pelo menos um departamento para receber a programação.", "error")
        return redirect(url_for("communication_center", year=year, month=month))

    created_count = 0
    created_events = 0
    skipped_departments = []
    for department_id in recipient_ids:
        dept = db.execute("SELECT * FROM departments WHERE id = ? AND active = 1", (department_id,)).fetchone()
        if not dept:
            continue
        period = f"{year:04d}-{month:02d}"
        event_rows = db.execute(
            "SELECT e.*, u.involvement AS area_involvement, u.details AS area_details, "
            "u.area_responsible AS area_responsible, u.due_date AS due_date "
            "FROM events e LEFT JOIN event_department_updates u ON u.event_id=e.id AND u.department_id=? "
            "WHERE substr(e.start_date,1,7)=? AND e.status IN ('aprovado','em_planejamento') "
            "ORDER BY e.start_date, e.start_time, e.title",
            (department_id, period),
        ).fetchall()
        relevant = [dict(r) for r in event_rows if (r["area_involvement"] or "conhecimento") != "excluir"]
        if not relevant:
            skipped_departments.append(dept["name"])
            continue
        for item in relevant:
            item["involvement"] = item.get("area_involvement") or "conhecimento"
            item["details"] = item.get("area_details") or ""
            item["response_url"] = ""
            if item["involvement"] == "acao":
                area_row = db.execute("SELECT id,response_token FROM event_department_updates WHERE event_id=? AND department_id=?", (item["id"], department_id)).fetchone()
                if area_row:
                    response_token = area_row["response_token"] or secrets.token_urlsafe(32)
                    db.execute("UPDATE event_department_updates SET response_token=?, response_status=CASE WHEN response_status IS NULL OR response_status='' THEN 'aguardando' ELSE response_status END WHERE id=?", (response_token, area_row["id"]))
                    base_url = _public_app_base_url()
                    if base_url:
                        item["response_url"] = base_url + url_for("department_response", token=response_token)
        subject, body = build_monthly_message(dept, year, month, relevant)
        cursor = db.execute(
            "INSERT INTO communication_batches(department_id, period_year, period_month, subject, body, recipient_email, status, created_by, created_at) "
            "VALUES (?, ?, ?, ?, ?, ?, 'rascunho', ?, ?)",
            (department_id, year, month, subject, body, dept["recipient_email"], session.get("user_email"), now_iso()),
        )
        batch_id = cursor.lastrowid
        for item in relevant:
            db.execute(
                "INSERT INTO communication_batch_events(batch_id, event_id, event_title, event_date, event_time, event_venue, involvement, details) "
                "VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                (batch_id, item["id"], item["title"], item["start_date"],
                 (f"{item['start_time']} às {item['end_time']}" if item.get("start_time") and item.get("end_time") else item.get("start_time")),
                 item["venue"], item["involvement"], item["details"]),
            )
        created_count += 1
        created_events += len(relevant)

    db.commit()
    if created_count:
        message = f"Foram criadas {created_count} prévia(s), totalizando {created_events} evento(s). Revise e envie na mesma página."
        if skipped_departments:
            message += " Sem eventos incluídos: " + ", ".join(skipped_departments) + "."
        flash(message, "success")
    else:
        flash("Não foi possível gerar prévias: não há eventos incluíveis nos departamentos selecionados para esse período.", "error")

    edit_department_id = request.form.get("edit_department_id", type=int)
    query = {"year": year, "month": month, "show_drafts": 1}
    if edit_department_id:
        query["edit_department_id"] = edit_department_id
    for dept_id in recipient_ids:
        query.setdefault("recipient_department_id", [])
        query["recipient_department_id"].append(dept_id)
    return redirect(url_for("communication_center", **query) + "#rascunhos-envio")


@app.get("/comunicacao/comunicado/<int:batch_id>")
@login_required
def communication_batch_detail(batch_id):
    db = get_db()
    batch = db.execute("SELECT b.*, d.name AS department_name FROM communication_batches b JOIN departments d ON d.id=b.department_id WHERE b.id=?", (batch_id,)).fetchone()
    if not batch:
        abort(404)
    items = db.execute("SELECT * FROM communication_batch_events WHERE batch_id = ? ORDER BY event_date, event_time, event_title", (batch_id,)).fetchall()
    return render_template("communication_batch.html", batch=batch, items=items, smtp_configured=smtp_is_configured())


def _communication_batch_redirect(batch):
    """Return to the monthly communication workspace when an inline action asks for it."""
    if request.form.get("return_to_center") == "1":
        return redirect(url_for(
            "communication_center", year=batch["period_year"], month=batch["period_month"],
            edit_department_id=batch["department_id"], show_drafts=1,
        ) + "#rascunhos-envio")
    return redirect(url_for("communication_batch_detail", batch_id=batch["id"]))


def _save_inline_batch_fields(db, batch):
    """Save editable draft fields submitted by the one-page communication workspace."""
    if "subject" not in request.form and "body" not in request.form and "recipient_email" not in request.form:
        return batch, True
    if batch["status"] != "rascunho":
        flash("Somente comunicados em rascunho podem ser alterados.", "error")
        return batch, False
    subject = request.form.get("subject", "").strip()[:250]
    body = request.form.get("body", "").strip()[:30000]
    recipient = request.form.get("recipient_email", "").strip().lower() or None
    if not subject or not body:
        flash("Informe o assunto e o conteúdo do comunicado.", "error")
        return batch, False
    if "\r" in subject or "\n" in subject or (recipient and ("\r" in recipient or "\n" in recipient)):
        flash("Não use quebras de linha no assunto ou no destinatário.", "error")
        return batch, False
    if recipient and ("@" not in recipient or "." not in recipient.split("@")[-1]):
        flash("O e-mail de destinatário não é válido.", "error")
        return batch, False
    db.execute("UPDATE communication_batches SET subject=?, body=?, recipient_email=? WHERE id=?", (subject, body, recipient, batch["id"]))
    db.commit()
    return db.execute("SELECT * FROM communication_batches WHERE id=?", (batch["id"],)).fetchone(), True


@app.post("/comunicacao/comunicado/<int:batch_id>/atualizar")
@login_required
def update_communication_batch(batch_id):
    db = get_db()
    batch = db.execute("SELECT * FROM communication_batches WHERE id = ?", (batch_id,)).fetchone()
    if not batch:
        abort(404)
    if batch["status"] != "rascunho":
        flash("Somente comunicados em rascunho podem ser editados.", "error")
    else:
        subject = request.form.get("subject", "").strip()[:250]
        body = request.form.get("body", "").strip()[:30000]
        recipient = request.form.get("recipient_email", "").strip().lower() or None
        if not subject or not body:
            flash("Informe o assunto e o conteúdo do comunicado.", "error")
        elif "\r" in subject or "\n" in subject or (recipient and ("\r" in recipient or "\n" in recipient)):
            flash("Não use quebras de linha no assunto ou no destinatário.", "error")
        elif recipient and ("@" not in recipient or "." not in recipient.split("@")[-1]):
            flash("O e-mail de destinatário não é válido.", "error")
        else:
            db.execute("UPDATE communication_batches SET subject=?, body=?, recipient_email=? WHERE id=?", (subject, body, recipient, batch_id))
            db.commit()
            flash("Prévia e destinatário salvos.", "success")
    return _communication_batch_redirect(batch)


def _mark_batch_sent(db, batch_id: int, method: str) -> None:
    db.execute("UPDATE communication_batches SET status='enviado', sent_at=?, sent_by=?, send_method=?, error_message=NULL WHERE id=?",
               (now_iso(), session.get("user_email"), method, batch_id))
    db.commit()


@app.post("/comunicacao/comunicado/<int:batch_id>/enviar")
@login_required
def send_communication_batch(batch_id):
    db = get_db()
    batch = db.execute("SELECT * FROM communication_batches WHERE id=?", (batch_id,)).fetchone()
    if not batch:
        abort(404)
    if batch["status"] != "rascunho":
        flash("Este comunicado já foi marcado como enviado e não pode ser reenviado por esta ação.", "error")
        return _communication_batch_redirect(batch)
    batch, fields_ok = _save_inline_batch_fields(db, batch)
    if not fields_ok:
        return _communication_batch_redirect(batch)
    recipient = (batch["recipient_email"] or "").strip()
    cfg = smtp_settings()
    if not recipient or not smtp_is_configured():
        flash("O envio automático ainda não está configurado. Cadastre o e-mail da área e configure SMTP no arquivo .env; ou envie pelo e-mail institucional e depois registre o envio manual.", "error")
        return _communication_batch_redirect(batch)
    if any(ch in recipient or ch in (batch["subject"] or "") for ch in ["\r", "\n"]):
        flash("O destinatário ou assunto contém caracteres não permitidos.", "error")
        return _communication_batch_redirect(batch)
    msg = EmailMessage()
    msg["Subject"] = batch["subject"]
    msg["From"] = cfg["from_email"]
    msg["To"] = recipient
    msg.set_content(batch["body"], cte="quoted-printable")
    try:
        with smtplib.SMTP(cfg["host"], cfg["port"], timeout=20) as server:
            if cfg["use_tls"]:
                server.starttls()
            if cfg["user"]:
                server.login(cfg["user"], cfg["password"])
            server.send_message(msg)
        _mark_batch_sent(db, batch_id, "smtp")
        flash(f"Comunicado enviado para {recipient} e registrado no histórico.", "success")
    except Exception as exc:
        safe_error = f"Falha no envio SMTP: {type(exc).__name__}"
        db.execute("UPDATE communication_batches SET error_message=? WHERE id=?", (safe_error, batch_id))
        db.commit()
        flash("Não foi possível enviar o comunicado. Confira a configuração com a TI; o rascunho foi preservado.", "error")
    return _communication_batch_redirect(batch)


@app.post("/comunicacao/comunicado/<int:batch_id>/registrar-envio-manual")
@login_required
def register_manual_communication_batch(batch_id):
    db = get_db()
    batch = db.execute("SELECT * FROM communication_batches WHERE id=?", (batch_id,)).fetchone()
    if not batch:
        abort(404)
    if batch["status"] != "rascunho":
        flash("Este comunicado já está registrado como enviado.", "error")
    else:
        batch, fields_ok = _save_inline_batch_fields(db, batch)
        if not fields_ok:
            return _communication_batch_redirect(batch)
        if not (batch["recipient_email"] or "").strip():
            flash("Antes de registrar o envio manual, informe o e-mail destinatário.", "error")
        elif request.form.get("confirm_sent") != "yes":
            flash("Marque a confirmação de que o e-mail foi enviado fora do sistema.", "error")
        else:
            _mark_batch_sent(db, batch_id, "manual")
            flash("Envio manual registrado no histórico.", "success")
    return _communication_batch_redirect(batch)


@app.post("/comunicacao/importar-agenda")
@login_required
def preview_agenda_import():
    if session.get("user_role") != "admin":
        abort(404)
    upload = request.files.get("agenda_file")
    if not upload or not upload.filename:
        flash("Selecione a planilha .xlsx para importar.", "error")
        return redirect(url_for("communication_center"))
    filename = Path(upload.filename).name
    if not filename.lower().endswith(".xlsx"):
        flash("Envie um arquivo .xlsx do Excel.", "error")
        return redirect(url_for("communication_center"))
    try:
        payload = parse_agenda_workbook(upload.read(), filename)
    except Exception as exc:
        flash(f"Não foi possível ler a planilha. Confira se o arquivo é um .xlsx válido ({type(exc).__name__}).", "error")
        return redirect(url_for("communication_center"))
    if not payload["events"]:
        flash(f"A planilha não contém eventos com data a partir de {payload['today_filter']} para importar.", "error")
        return redirect(url_for("communication_center"))
    _cleanup_import_previews()
    token = secrets.token_urlsafe(24)
    preview_dir = INSTANCE_DIR / "import_previews"
    preview_dir.mkdir(exist_ok=True)
    (preview_dir / f"{token}.json").write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return render_template("agenda_import_preview.html", payload=payload, token=token)


@app.post("/comunicacao/importar-agenda/confirmar")
@login_required
def confirm_agenda_import():
    if session.get("user_role") != "admin":
        abort(404)
    token = request.form.get("token", "")
    if not re.fullmatch(r"[A-Za-z0-9_-]{20,80}", token):
        abort(400, description="Identificador de prévia inválido.")
    preview_path = INSTANCE_DIR / "import_previews" / f"{token}.json"
    if not preview_path.exists():
        flash("A prévia expirou ou já foi utilizada. Envie a planilha novamente.", "error")
        return redirect(url_for("communication_center"))
    try:
        payload = json.loads(preview_path.read_text(encoding="utf-8"))
        inserted, skipped, links, tasks = import_agenda_payload(payload)
        preview_path.unlink(missing_ok=True)
        flash(f"Importação concluída: {inserted} evento(s) inserido(s), {skipped} duplicado(s) ignorado(s), {links} vínculos com áreas e {tasks} tarefas de checklist criadas.", "success")
        imported_ids = set(payload.get("_imported_ids") or [])
        if imported_ids:
            pool = get_db().execute("SELECT * FROM events WHERE status<>'cancelado' AND start_date>=? ORDER BY start_date,start_time",
                                    (date.today().isoformat(),)).fetchall()
            clashes = [c for c in _find_schedule_conflicts(pool) if {c["event_a_id"], c["event_b_id"]} & imported_ids]
            if clashes:
                flash(f"Atenção: {len(clashes)} possível(is) conflito(s) de local/horário envolvendo eventos importados. Confira a Agenda Institucional.", "warning")
    except Exception as exc:
        get_db().rollback()
        flash(f"A importação não foi concluída; a prévia foi preservada. Erro: {type(exc).__name__}.", "error")
    return redirect(url_for("communication_center"))


__all__ = [name for name in dict(globals()) if not name.startswith("__")]
